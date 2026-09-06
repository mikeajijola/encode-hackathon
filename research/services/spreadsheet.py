"""SpreadsheetBench service wiring; golden workbooks are never accepted or read."""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Mapping
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from adapters.spreadsheet import (
    KIND, VERSION, SpreadsheetCapability, WorkbookSnapshots, expand_selector_scopes,
    manifests, parse_selector, register_spreadsheet_capabilities, scope_for,
)
from experiment.runner import EvaluationResult, ExecutionResult, Task, TaskRuntime
from sb import answer_cells
from fulfilment import (
    Broker, CapabilityRequest, Contract, ContractValidator, DesiredAssertion,
    Discrepancy, EvalResult, EvalSpec, EvalStatus, EvidenceStore, Observation,
    Scope, derive_discrepancies,
    DeterministicContractCompiler, EvaluatorRegistry, FulfilmentAgent, decide_completion,
    Transition, TransitionKind,
)
from fulfilment.models import FrozenDict


CONTRACT_SCHEMA_VERSION = "2.0.0"
CONTRACT_PROPOSAL_FIELDS = {
    "schema_version", "desired_state", "constraints", "invariants",
    "evaluator_intents", "required_capabilities",
}
ASSERTION_FIELDS = {"id", "property", "description", "output"}
OUTPUT_FIELDS = {"type", "shape", "uncertainty_allowed"}
EVALUATOR_INTENT_FIELDS = {"id", "assertion_id", "evaluator", "purpose"}
CAPABILITY_REF_FIELDS = {"name", "version"}
OUTPUT_TYPES = {"number", "text", "boolean", "date", "formula", "mixed", "any"}
OUTPUT_SHAPES = {"scalar", "range"}
STATE_PROPERTIES = {"computed_value", "formula_result", "transformed_values", "table_state", "visual_state"}
AVAILABLE_EVALUATORS = {
    "artifact-valid", "preservation", "nonblank", "type", "output-shape",
    "formula-errors", "semantic-independent", "render",
}
PROCEDURAL_PREFIXES = ("first ", "then ", "next ", "step ", "click ", "open ", "write ", "copy ")
TRANSITION_FIELDS = {"capability", "inputs", "rationale"}
TRANSITION_CAPABILITY_FIELDS = {"name", "version"}
MAX_LITERAL_WRITES = 400
MAX_INSPECT_CELLS = 400


@dataclass
class _Session:
    contract: Contract
    broker: Broker
    snapshots: WorkbookSnapshots
    evidence: EvidenceStore
    initial_cells: dict[str, Any]
    discrepancies: tuple[Discrepancy, ...] = ()
    capability_observations: list[dict[str, Any]] = field(default_factory=list)
    first_mutation_path: Path | None = None
    mutation_paths: list[Path] = field(default_factory=list)


class _RuntimeCapability:
    def __init__(self, runtime, handler, session):
        self.runtime, self.handler, self.session = runtime, handler, session
    def invoke(self, request):
        self.runtime.action(request.capability_name, {"discrepancy_ids": request.discrepancy_ids})
        result = self.handler.invoke(request)
        self.session.capability_observations.append({
            "capability": request.capability_name, "succeeded": result.succeeded,
            "output": dict(result.output), "error": result.error,
            "artifact_sha256": result.provenance.get("artifact_sha256"),
        })
        if result.succeeded and result.actual_mutation_scope and self.session.first_mutation_path is None:
            root = self.runtime.event_path.parent.parent
            checkpoint = root / "checkpoints" / f"{self.runtime.event_path.stem}.first_mutation.xlsx"
            checkpoint.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(request.artifact_id, checkpoint)
            self.session.first_mutation_path = checkpoint
            self.runtime.event("first_mutation_checkpoint", {
                "path": str(checkpoint.relative_to(root)), "artifact_hash": _file_hash(checkpoint),
                "capability": request.capability_name, "request_id": request.id,
            })
        if result.succeeded and result.actual_mutation_scope:
            root = self.runtime.event_path.parent.parent
            number = len(self.session.mutation_paths) + 1
            checkpoint = root / "checkpoints" / self.runtime.event_path.stem / f"mutation-{number:02d}.xlsx"
            checkpoint.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(request.artifact_id, checkpoint)
            self.session.mutation_paths.append(checkpoint)
            self.runtime.event("mutation_checkpoint", {
                "mutation_number": number, "path": str(checkpoint.relative_to(root)),
                "artifact_hash": _file_hash(checkpoint), "capability": request.capability_name,
                "request_id": request.id, "discrepancy_ids": list(request.discrepancy_ids),
                "usage": self.runtime.usage(),
            })
        return result


class _RuntimeEvidenceStore(EvidenceStore):
    """Hash-chain remains authoritative; runtime stream gets a searchable mirror."""
    def __init__(self, path, runtime): self.runtime = runtime; super().__init__(path)
    def append(self, event_type, payload):
        event = super().append(event_type, payload)
        mirrored = {"evidence_event_id": event.id, **dict(payload),
                    "usage": self.runtime.usage()}
        if event_type == "capability_result":
            mirrored["actual_scope"] = list(payload.get("actual_mutation_scope", ()))
        self.runtime.event(event_type, mirrored)
        return event


class _SpreadsheetObserver:
    def __init__(self, task, artifact):
        self.task, self.artifact, self.count = task, artifact, 0

    def observe(self, scopes=()):
        self.count += 1
        selectors = _answer_selectors(self.task.context, self.artifact)
        hidden = bool(self.task.context.get("test_omit_target_once")) and self.count == 1 and not scopes
        facts = [] if hidden else _selected_cells(self.artifact, selectors)
        inspected = () if hidden else tuple(scope_for(item) for item in selectors)
        omitted = tuple(scope_for(item) for item in selectors) if hidden else (scope_for("__omitted__!A1"),)
        return Observation(str(uuid4()), str(self.artifact), KIND, _file_hash(self.artifact),
                           {"cells": facts}, {}, inspected, omitted)


class _SpreadsheetEvaluator:
    def __init__(self, services, task, artifact, session, runtime):
        self.services, self.task, self.artifact = services, task, artifact
        self.session, self.runtime, self.latest = session, runtime, []

    def evaluate(self, spec, observation, contract):
        facts = list(observation.facts.get("cells", ()))
        values = [item["value"] for item in facts]
        if spec.evaluator == "artifact-valid":
            valid, error = _valid(self.artifact)
            result = _eval(spec.id, observation, valid, error or "workbook loads", kind="artifact_invalid")
        elif spec.evaluator == "preservation":
            current = _all_cells(self.artifact)
            allowed = {scope.resource for scope in contract.authorized_mutation_scopes}
            changed = sorted(key for key in self.session.initial_cells.keys() | current.keys()
                             if key not in allowed and self.session.initial_cells.get(key) != current.get(key))
            result = _eval(spec.id, observation, not changed, "outside scope preserved",
                           observed=changed, kind="constraint_violation")
        elif not facts:
            result = EvalResult(spec.id, observation.id, EvalStatus.FAIL, "target facts unobserved",
                                details={"knowledge_gap": True, "expected": "observed target facts",
                                         "confidence": 1.0})
        elif spec.evaluator == "nonblank":
            result = _eval(spec.id, observation, all(value not in (None, "") for value in values),
                           "answers nonblank", observed=values)
        elif spec.evaluator == "type":
            expected = spec.parameters.get("expected_type", "any")
            result = _eval(spec.id, observation, expected in ("any", "mixed") or
                           all(_type_name(value) == expected for value in values), "answer types",
                           expected=expected, observed=[_type_name(value) for value in values])
        elif spec.evaluator == "output-shape":
            expected = spec.parameters["expected_shape"]; observed = "scalar" if len(values) == 1 else "range"
            result = _eval(spec.id, observation, expected == observed, "answer shape",
                           expected=expected, observed=observed)
        elif spec.evaluator == "formula-errors":
            current = _all_cells(self.artifact)
            initial_errors = {key for key, value in self.session.initial_cells.items()
                              if isinstance(value, str) and value.startswith("#")}
            errors = sorted(key for key, value in current.items() if isinstance(value, str) and
                            value.startswith("#") and key not in initial_errors)
            result = _eval(spec.id, observation, not errors, "no new formula errors", observed=errors)
        elif spec.evaluator == "semantic-independent":
            if "independent_expected" in self.task.context:
                expected = self.task.context["independent_expected"]
                actual = values[0] if len(values) == 1 else values
                result = _eval(spec.id, observation, actual == expected, "independent computation",
                               expected=expected, observed=actual)
            elif any(value in (None, "") for value in values):
                result = _eval(spec.id, observation, False, "semantic evaluation deferred until target exists",
                               expected="intent-satisfying target state", observed=values,
                               likely_causes=["target state is missing"], confidence=1.0)
            else:
                result = self.services._semantic_model_eval(self.task, contract, observation, facts, self.runtime)
        elif spec.evaluator == "render":
            selector = _answer_selectors(self.task.context, self.artifact)[0]
            request = CapabilityRequest(str(uuid4()), "render_range", VERSION, str(self.artifact), KIND,
                {"selector": selector, "output_dir": str(self.artifact.parent / "renders")}, (), ())
            rendered = self.session.broker.invoke(contract, request)
            result = _eval(spec.id, observation, rendered.succeeded, rendered.error or "render captured",
                           kind="evaluation_uncertainty")
        else:
            result = EvalResult(spec.id, observation.id, EvalStatus.ERROR, "unsupported spreadsheet evaluator")
        self.latest = [item for item in self.latest if item.eval_id != result.eval_id] + [result]
        self.runtime.event("eval_result", _eval_mapping(result))
        return result


class _SpreadsheetPlanner:
    def __init__(self, services, task, artifact, runtime, session):
        self.services, self.task, self.artifact, self.runtime, self.session = services, task, artifact, runtime, session

    def plan(self, contract, observation, discrepancies, capabilities):
        ids = tuple(item.id for item in discrepancies)
        if any(item.kind.value == "knowledge_discrepancy" for item in discrepancies):
            scopes = tuple(scope for item in discrepancies for scope in item.permissible_mutation_scope)
            return Transition(str(uuid4()), TransitionKind.OBSERVE, ids, "observe missing target facts", scopes)
        if any(item.kind.value == "evaluation_uncertainty" for item in discrepancies):
            return None
        proposal = self.services._action_proposal(self.task, _contract_mapping_existing(contract),
                                                  self.runtime, discrepancies, session=self.session)
        name = proposal["capability"]["name"]
        version = proposal["capability"]["version"]
        if not any(item.name == name and item.version == version for item in capabilities):
            return None
        inputs = dict(proposal["inputs"])
        if name == "write_cells":
            requested = tuple(scope_for(item["selector"]) for item in inputs["writes"])
        elif name == "copy_or_fill_formula":
            requested = expand_selector_scopes(inputs["target"])
        elif name == "recalculate":
            requested = contract.authorized_mutation_scopes
        else:
            requested = ()
        allowed = {scope.resource for scope in contract.authorized_mutation_scopes}
        if any(scope.resource not in allowed for scope in requested):
            self.runtime.event("scope_violation_rejected", {"requested": [s.resource for s in requested]})
            return None
        if name == "render_range":
            inputs["output_dir"] = str(self.artifact.parent / "renders")
        request = CapabilityRequest(str(uuid4()), name, version, str(self.artifact), KIND,
                                    inputs, requested, ids)
        return Transition(str(uuid4()), TransitionKind.CAPABILITY, ids,
                          proposal["rationale"], capability_request=request)


class _SpreadsheetLifecycle:
    def __init__(self, artifact, runtime): self.artifact, self.runtime = artifact, runtime
    def rollback(self, reason): self.runtime.event("capability_rollback", {"reason": reason})
    def select_result(self, observation, status):
        return {"path": str(self.artifact), "hash": _file_hash(self.artifact), "status": status}


class SpreadsheetServices:
    """Implements runner Services while delegating all XLSX effects to the adapter."""

    FORBIDDEN_CONTEXT_KEYS = {"golden", "golden_xlsx", "answer", "expected_answer"}

    def __init__(self):
        self._sessions: dict[str, _Session] = {}

    def compile_contract(self, task: Task, runtime: TaskRuntime) -> Mapping[str, Any]:
        forbidden = self.FORBIDDEN_CONTEXT_KEYS & set(task.context)
        if forbidden:
            raise ValueError(f"runtime context contains forbidden golden-like keys: {sorted(forbidden)}")
        selectors = _answer_selectors(task.context, task.artifact)
        compiler_input = {
            "contract_schema_version": CONTRACT_SCHEMA_VERSION,
            "instruction": (
                "Return exactly one JSON object, optionally enclosed by one ```json code fence. "
                "Describe required state/properties, never actions or steps. Use exactly the declared "
                "schema fields and exactly ONE desired_state assertion. The assertion output.shape MUST "
                "equal required_output_shape. Include semantic-independent in evaluator_intents and also "
                "render when visual_intent is true. Cite only supplied capability name/version pairs and "
                "evaluator identifiers. Do not guess hidden expected values."
            ),
            "schema": {
                "schema_version": CONTRACT_SCHEMA_VERSION,
                "desired_state": {"assertions": [{"id": "string", "property": sorted(STATE_PROPERTIES),
                    "description": "state assertion", "output": {"type": sorted(OUTPUT_TYPES),
                    "shape": sorted(OUTPUT_SHAPES), "uncertainty_allowed": "boolean"}}]},
                "constraints": ["state constraint"], "invariants": ["state invariant"],
                "evaluator_intents": [{"id": "string", "assertion_id": "string",
                    "evaluator": sorted(AVAILABLE_EVALUATORS), "purpose": "state property to test"}],
                "required_capabilities": [{"name": "manifest name", "version": "manifest version"}],
            },
            "intent": task.intent,
            "required_output_shape": "scalar" if _selector_cell_count(selectors) == 1 else "range",
            "visual_intent": bool(task.context.get("visual_intent")),
            "observed_context": _model_context(task.context),
            "capability_manifests": task.capability_manifests,
        }
        reply = runtime.complete(
            json.dumps(compiler_input, default=str),
            purpose="contract",
        )
        proposal = _validate_contract_proposal(_strict_contract_reply(reply.text), task, selectors)
        proposed = proposal["desired_state"]["assertions"][0]
        assertion_id = proposed["id"]
        assertion = DesiredAssertion(assertion_id, proposed["description"], tuple(scope_for(s) for s in selectors))
        purpose_by_evaluator = {item["evaluator"]: item["purpose"] for item in proposal["evaluator_intents"]}
        output = proposed["output"]
        evals = [
            EvalSpec("artifact-valid", assertion_id, "artifact-valid"),
            EvalSpec("preservation", assertion_id, "preservation"),
            EvalSpec("nonblank", assertion_id, "nonblank"),
            EvalSpec("type", assertion_id, "type", parameters={"expected_type": output["type"]}),
            EvalSpec("output-shape", assertion_id, "output-shape", parameters={"expected_shape": output["shape"]}),
            EvalSpec("formula-errors", assertion_id, "formula-errors"),
            EvalSpec("semantic", assertion_id, "semantic-independent", parameters={
                "property": proposed["property"], "purpose": purpose_by_evaluator["semantic-independent"],
                "uncertainty_allowed": output["uncertainty_allowed"],
            }),
        ]
        if task.context.get("visual_intent"):
            evals.append(EvalSpec("visual", assertion_id, "render",
                                  parameters={"purpose": purpose_by_evaluator["render"]}))
        policy_constraint = "preserve cells outside authorized answer selectors"
        policy_invariant = "workbook remains structurally valid"
        contract = Contract(
            f"contract-{task.id}", 2, task.intent, (assertion,),
            tuple(dict.fromkeys((policy_constraint, *proposal["constraints"]))),
            tuple(dict.fromkeys((policy_invariant, *proposal["invariants"]))),
            tuple(evals), tuple(scope_for(s) for s in selectors),
        )
        ContractValidator().validate(contract)
        return _contract_mapping(contract, task.context, proposed, proposal["required_capabilities"])

    def execute_once(self, task: Task, contract: Mapping[str, Any] | None, destination: Path,
                     runtime: TaskRuntime) -> ExecutionResult:
        shutil.copy2(task.artifact, destination)
        session = self._session(task, contract or self._direct_contract(task), destination, runtime,
                                inspect=contract is not None, record_contract=contract is not None)
        proposal = self._action_proposal(task, contract, runtime, discrepancies=(), session=session)
        result = self._apply_proposal(task, destination, runtime, session, proposal, ())
        return ExecutionResult(destination, "ok" if result else "action_failed", {"broker_evidence": str(session.evidence.path)})

    def evaluate_once(self, task: Task, contract: Mapping[str, Any], artifact: Path,
                      runtime: TaskRuntime) -> EvaluationResult:
        session = self._session(task, contract, artifact, runtime, inspect=True)
        results, discrepancies = self._evaluate(task, artifact, session, runtime)
        session.discrepancies = discrepancies
        by_id = {result.eval_id: result for result in results}
        decision = decide_completion(
            session.contract, results, session.evidence.records(),
            desired_state_satisfied=not discrepancies,
            constraints_preserved=by_id.get("preservation") is not None and by_id["preservation"].passed,
            invariants_hold=all(by_id.get(key) is not None and by_id[key].passed
                                for key in ("artifact-valid", "formula-errors")),
            artifact_valid=by_id.get("artifact-valid") is not None and by_id["artifact-valid"].passed,
        )
        status = ("pass" if decision.fulfilled else
                  "uncertain" if any(r.status is EvalStatus.UNCERTAIN for r in results) else
                  "completion_gate_failed")
        details = {"evals": [_eval_mapping(r) for r in results],
                   "discrepancies": [_discrepancy_mapping(d) for d in discrepancies],
                   "completion": asdict(decision)}
        session.evidence.append("termination_decision", {
            "fulfilled": decision.fulfilled, "reason": status, "completion": asdict(decision),
            "artifact": {"path": str(artifact), "hash": _file_hash(artifact)},
        })
        runtime.event("completion_decision", {"fulfilled": decision.fulfilled,
                                               "failed_conditions": decision.failed_conditions})
        return EvaluationResult(decision.fulfilled, status, details)

    def reconcile(self, task: Task, contract: Mapping[str, Any], destination: Path,
                  runtime: TaskRuntime) -> tuple[ExecutionResult, EvaluationResult]:
        shutil.copy2(task.artifact, destination)
        session = self._session(task, contract, destination, runtime, inspect=False, record_contract=False,
                                runtime_capabilities=True)
        observer = _SpreadsheetObserver(task, destination)
        domain_evaluator = _SpreadsheetEvaluator(self, task, destination, session, runtime)
        registry = EvaluatorRegistry(session.evidence)
        for name in {spec.evaluator for spec in session.contract.evals}:
            registry.register(name, domain_evaluator)
        agent = FulfilmentAgent(
            DeterministicContractCompiler({task.intent: session.contract}), ContractValidator(), observer,
            registry, _SpreadsheetPlanner(self, task, destination, runtime, session), session.broker,
            session.evidence, _SpreadsheetLifecycle(destination, runtime),
            max_iterations=runtime.config.max_actions, max_repeated_transition=1,
        )
        result = agent.run(task.intent, observer.observe(()))
        latest = domain_evaluator.latest
        evaluation = EvaluationResult(result.fulfilled, "pass" if result.fulfilled else result.reason, {
            "iteration": result.iterations, "evals": [_eval_mapping(item) for item in latest],
            "discrepancies": [_discrepancy_mapping(item) for item in result.open_discrepancies],
        })
        runtime.event("terminal_evaluation", asdict(evaluation))
        status = "fulfilled" if result.fulfilled else f"unfulfilled:{result.reason}"
        return ExecutionResult(destination, status, {"iterations": result.iterations,
            "broker_evidence": str(session.evidence.path),
            "first_mutation_artifact": str(session.first_mutation_path) if session.first_mutation_path else None,
            "mutation_artifacts": [str(path) for path in session.mutation_paths]}), evaluation

    def _session(self, task, contract_mapping, artifact, runtime, *, inspect=False, record_contract=True,
                 runtime_capabilities=False):
        if task.id in self._sessions:
            return self._sessions[task.id]
        contract = _contract_from_mapping(contract_mapping)
        evidence_path = runtime.event_path.with_name(f"{task.id}.broker.jsonl")
        evidence = _RuntimeEvidenceStore(evidence_path, runtime) if runtime_capabilities else EvidenceStore(evidence_path)
        snapshots = WorkbookSnapshots()
        broker = Broker(evidence, snapshots)
        if not runtime_capabilities:
            register_spreadsheet_capabilities(broker)
        session = _Session(contract, broker, snapshots, evidence, _all_cells(artifact))
        if runtime_capabilities:
            for manifest in manifests():
                broker.register(manifest, _RuntimeCapability(runtime, SpreadsheetCapability(manifest.name), session))
        self._sessions[task.id] = session
        if record_contract:
            evidence.append("accepted_contract", {"contract_id": contract.id, "version": contract.version,
                                                   "intent": contract.intent})
        if inspect:
            selectors = _answer_selectors(task.context, artifact)
            runtime.action("inspect_workbook", {"selectors": selectors})
            request = CapabilityRequest(str(uuid4()), "inspect_workbook", VERSION, str(artifact), KIND,
                                        {"selectors": selectors}, (), ())
            observed = broker.invoke(contract, request)
            runtime.event("bounded_observation", dict(observed.output["observation"]))
        return session

    def _direct_contract(self, task):
        selectors = _answer_selectors(task.context, task.artifact)
        shape = "scalar" if _selector_cell_count(selectors) == 1 else "range"
        proposed = {"id": "direct-state", "property": "transformed_values",
                    "description": "direct action boundary",
                    "output": {"type": task.context.get("expected_type") or "any", "shape": shape,
                               "uncertainty_allowed": True}}
        assertion = DesiredAssertion("direct-state", proposed["description"], tuple(scope_for(s) for s in selectors))
        evals = _policy_evals(assertion.id, proposed, "direct boundary", visual=False)
        contract = Contract(f"direct-{task.id}", 2, task.intent, (assertion,),
            ("preserve cells outside authorized answer selectors",),
            ("workbook remains structurally valid",), evals, tuple(scope_for(s) for s in selectors))
        return _contract_mapping(contract, task.context, proposed, [])

    def _action_proposal(self, task, contract, runtime, discrepancies, session=None):
        prompt = {
            "role": "bounded_transition_planner",
            "instruction": (
                "Return exactly one JSON object, optionally enclosed by one ```json code fence. It must have "
                "exactly fields capability, inputs, and rationale. capability MUST be an object with exactly "
                "name and version, for example {\"name\":\"write_cells\",\"version\":\"1.0.0\"}; never "
                "return capability as a string. Select an exact supplied name/version. Never emit code or "
                "execute commands. Prefer copy_or_fill_formula for large repeated formula ranges instead of "
                "literal cell writes."
            ),
            "required_output_schema": {
                "capability": {"name": "supplied capability name", "version": "supplied capability version"},
                "inputs": "the exact inputs object for that capability", "rationale": "non-empty string",
            },
            "transition_schemas": {
                "write_cells": {"inputs": {"writes": [{"selector": "Sheet!A1", "value": "JSON value"}]},
                                "max_writes": MAX_LITERAL_WRITES},
                "copy_or_fill_formula": {"inputs": {"source": "Sheet!A1", "target": "Sheet!A2:A10"}},
                "inspect_workbook": {"inputs": {"selectors": ["Sheet!A1:B10"]},
                                     "max_cells": MAX_INSPECT_CELLS},
                "recalculate": {"inputs": {}}, "validate_workbook": {"inputs": {}},
                "render_range": {"inputs": {"selector": "Sheet!A1:B10"}},
            },
            "intent": task.intent, "answer_selectors": _answer_selectors(task.context, task.artifact),
            "observed_context": _model_context(task.context), "accepted_contract": contract,
            "discrepancies": [_discrepancy_mapping(d) for d in discrepancies],
            "capability_manifests": task.capability_manifests,
            "capability_observations": tuple(session.capability_observations) if session else (),
        }
        reply = runtime.complete(json.dumps(prompt, default=str), purpose="action_generation")
        return _validate_transition(_strict_contract_reply(reply.text), task)

    def _apply_proposal(self, task, destination, runtime, session, proposal, discrepancies):
        name = proposal["capability"]["name"]
        inputs = dict(proposal["inputs"])
        if name == "write_cells":
            requested = tuple(scope_for(item["selector"]) for item in inputs["writes"])
        elif name == "copy_or_fill_formula":
            requested = expand_selector_scopes(inputs["target"])
        elif name == "recalculate":
            requested = session.contract.authorized_mutation_scopes
        else:
            requested = ()
        allowed = {scope.resource for scope in session.contract.authorized_mutation_scopes}
        if any(scope.resource not in allowed for scope in requested):
            runtime.event("scope_violation_rejected", {"requested": [s.resource for s in requested]})
            return False
        discrepancy_ids = tuple(d.id for d in discrepancies) or (f"initial-{task.id}",)
        if name == "render_range":
            inputs["output_dir"] = str(destination.parent / "renders")
        runtime.action(name, {"discrepancy_ids": discrepancy_ids, "rationale": proposal["rationale"]})
        request = CapabilityRequest(str(uuid4()), name, proposal["capability"]["version"],
                                    str(destination), KIND, inputs, requested, discrepancy_ids)
        result = session.broker.invoke(session.contract, request)
        observation = {"capability": name, "succeeded": result.succeeded,
                       "output": dict(result.output), "error": result.error,
                       "artifact_sha256": result.provenance.get("artifact_sha256")}
        session.capability_observations.append(observation)
        runtime.event("capability_result", {"capability": name, "succeeded": result.succeeded,
                      "actual_scope": [s.resource for s in result.actual_mutation_scope],
                      "discrepancy_ids": discrepancy_ids, "output": dict(result.output),
                      "provenance": dict(result.provenance), "error": result.error})
        return result.succeeded

    def _evaluate(self, task, artifact, session, runtime):
        selectors = _answer_selectors(task.context, artifact)
        facts = _selected_cells(artifact, selectors)
        observation = Observation(str(uuid4()), str(artifact), KIND, _file_hash(artifact),
            {"cells": facts}, {}, tuple(scope_for(s) for s in selectors), (scope_for("__omitted__!A1"),))
        self._record_evaluation_evidence(session, observation, ())
        results = []
        valid, structural_error = _valid(artifact)
        results.append(_eval("artifact-valid", observation, valid, structural_error or "workbook loads",
                             kind="artifact_invalid"))
        current = _all_cells(artifact)
        allowed = {s.resource for s in session.contract.authorized_mutation_scopes}
        changed_outside = sorted(k for k in session.initial_cells.keys() | current.keys()
                                 if k not in allowed and session.initial_cells.get(k) != current.get(k))
        results.append(_eval("preservation", observation, not changed_outside, "outside scope preserved",
                             observed=changed_outside, kind="constraint_violation"))
        values = [item["value"] for item in facts]
        results.append(_eval("nonblank", observation, all(v not in (None, "") for v in values), "answers nonblank", observed=values))
        type_spec = next(spec for spec in session.contract.evals if spec.id == "type")
        expected_type = type_spec.parameters.get("expected_type", "any")
        type_ok = expected_type in ("any", "mixed") or all(_type_name(v) == expected_type for v in values)
        results.append(_eval("type", observation, type_ok, "answer types", expected=expected_type,
                             observed=[_type_name(v) for v in values]))
        shape_spec = next(spec for spec in session.contract.evals if spec.id == "output-shape")
        expected_shape = shape_spec.parameters["expected_shape"]
        observed_shape = "scalar" if len(values) == 1 else "range"
        results.append(_eval("output-shape", observation, expected_shape == observed_shape,
                             "answer shape", expected=expected_shape, observed=observed_shape))
        initial_errors = {key for key, value in session.initial_cells.items()
                          if isinstance(value, str) and value.startswith("#")}
        new_errors = sorted(key for key, value in current.items()
                            if isinstance(value, str) and value.startswith("#") and key not in initial_errors)
        results.append(_eval("formula-errors", observation, not new_errors, "no new formula errors", observed=new_errors))
        if "independent_expected" in task.context:
            expected = task.context["independent_expected"]
            actual = values[0] if len(values) == 1 else values
            results.append(_eval("semantic", observation, actual == expected, "independent computation", expected=expected, observed=actual))
        elif any(value in (None, "") for value in values):
            # Missing desired state is already actionable. Asking a semantic
            # judge before the first mutation adds cost and previously caused D
            # to terminate uncertain without ever attempting fulfilment.
            results.append(_eval("semantic", observation, False,
                "semantic evaluation deferred until target state exists",
                expected="intent-satisfying target state", observed=values,
                likely_causes=["target state is missing"], confidence=1.0))
        else:
            results.append(self._semantic_model_eval(task, session.contract, observation, facts, runtime))
        if task.context.get("visual_intent"):
            request = CapabilityRequest(str(uuid4()), "render_range", VERSION, str(artifact), KIND,
                {"selector": selectors[0], "output_dir": str(artifact.parent / "renders")}, (), ())
            rendered = session.broker.invoke(session.contract, request)
            results.append(_eval("visual", observation, rendered.succeeded, rendered.error or "render captured",
                                 kind="evaluation_uncertainty"))
        for result in results:
            runtime.event("eval_result", _eval_mapping(result))
            session.evidence.append("eval_result", {"eval_id": result.eval_id,
                "observation_id": result.observation_id, "status": result.status.value,
                "message": result.message, "details": dict(result.details)})
        discrepancies = derive_discrepancies(session.contract, results, session.discrepancies)
        for discrepancy in discrepancies:
            runtime.event("discrepancy", _discrepancy_mapping(discrepancy))
        return tuple(results), discrepancies

    def _record_evaluation_evidence(self, session, observation, _results):
        """Overridable seam used to prove missing evidence blocks completion."""
        session.evidence.append("observation", {
            "observation_id": observation.id, "artifact_id": observation.artifact_id,
            "artifact_hash": observation.artifact_hash,
            "inspected_scope": [scope.resource for scope in observation.inspected_scope],
            "omitted_scope": [scope.resource for scope in observation.omitted_scope],
        })

    def _semantic_model_eval(self, task, contract, observation, target_facts, runtime):
        """Isolated semantic verdict. Its raw exchange is never planner input."""
        prompt = {
            "role": "independent_semantic_evaluator",
            "instruction": (
                "Independently solve the user intent from the bounded source observation, then judge the current "
                "target facts. Do not trust or reconstruct the action response. Return JSON only with exactly: "
                "verdict (pass|fail|uncertain), expected_state, rationale, confidence (0..1)."
            ),
            "intent": task.intent,
            "accepted_contract": {
                "id": contract.id, "version": contract.version,
                "assertions": [a.description for a in contract.assertions],
                "constraints": list(contract.constraints), "invariants": list(contract.invariants),
                "evaluator_intents": [{"id": spec.id, "evaluator": spec.evaluator,
                    "parameters": dict(spec.parameters)} for spec in contract.evals],
            },
            "bounded_source_observation": _model_context(task.context).get("workbook_observation", {}),
            "current_target_facts": target_facts,
        }
        provider_request_id = None
        try:
            reply = runtime.complete(json.dumps(prompt, default=str), purpose="independent_evaluation")
            provider_request_id = reply.provider_request_id
            verdict = _strict_semantic_verdict(reply.text)
            status = {"pass": EvalStatus.PASS, "fail": EvalStatus.FAIL,
                      "uncertain": EvalStatus.UNCERTAIN}[verdict["verdict"]]
            message = verdict["rationale"]
            details = {"expected": verdict["expected_state"], "observed": target_facts,
                       "confidence": verdict["confidence"],
                       "likely_causes": [] if status is EvalStatus.PASS else [message]}
        except Exception as error:
            status = EvalStatus.UNCERTAIN
            message = f"independent evaluator unavailable or malformed: {type(error).__name__}: {error}"
            details = {"expected": "independently verified intent-satisfying state",
                       "observed": target_facts, "confidence": 0.0,
                       "likely_causes": ["semantic evaluator failure"]}
        runtime.event("semantic_evaluator_verdict", {
            "role": "independent_semantic_evaluator", "status": status.value,
            "provider_request_id": provider_request_id,
        })
        return EvalResult("semantic", observation.id, status, message, details=details)


def _eval(eval_id, observation, passed, message, **details):
    return EvalResult(eval_id, observation.id, EvalStatus.PASS if passed else EvalStatus.FAIL, message, details=details)


def _answer_selectors(context, artifact=None):
    sheet, position = context.get("answer_sheet"), context.get("answer_position")
    if sheet is not None and not isinstance(sheet, str):
        raise ValueError("answer_sheet must be a string or null")
    if not isinstance(position, str):
        raise ValueError("answer_position metadata is required")
    if artifact is None:
        if sheet is None:
            raise ValueError("artifact is required to resolve an active-sheet selector")
        return [f"{sheet}!{position}"]
    wb = load_workbook(artifact, read_only=True)
    synthetic = {"answer_sheet": sheet, "answer_position": position}
    selectors = [f"{resolved_sheet or wb.active.title}!{cell}"
                 for resolved_sheet, cell in answer_cells(synthetic, wb)]
    wb.close()
    return selectors


def _contract_mapping(contract, context, proposed_assertion, required_capabilities):
    value = {
        "id": contract.id, "version": contract.version, "intent": contract.intent,
        "desired_state": {"assertions": [{**proposed_assertion,
            "selectors": [s.resource for s in contract.assertions[0].target_scope]}]},
        "constraints": list(contract.constraints), "invariants": list(contract.invariants),
        "evals": [{"id": spec.id, "assertion_id": spec.assertion_id, "evaluator": spec.evaluator,
                   "severity": spec.severity, "parameters": dict(spec.parameters)} for spec in contract.evals],
        "authorized_mutation_scopes": [s.resource for s in contract.authorized_mutation_scopes],
        "required_capabilities": required_capabilities,
        "visual_intent": bool(context.get("visual_intent")),
        "has_independent_expected": "independent_expected" in context,
    }
    return _freeze_mapping(value)


def _contract_mapping_existing(contract):
    """Prompt-safe projection of the accepted immutable generic contract."""
    return {"id": contract.id, "version": contract.version, "intent": contract.intent,
            "assertions": [{"id": item.id, "description": item.description,
                            "target_scope": [scope.resource for scope in item.target_scope]}
                           for item in contract.assertions],
            "constraints": list(contract.constraints), "invariants": list(contract.invariants),
            "evals": [{"id": item.id, "evaluator": item.evaluator,
                       "parameters": dict(item.parameters)} for item in contract.evals]}


def _contract_from_mapping(value):
    proposed = value["desired_state"]["assertions"][0]
    scopes = tuple(Scope(resource) for resource in value["authorized_mutation_scopes"])
    assertion = DesiredAssertion(proposed["id"], proposed["description"], scopes)
    evals = tuple(EvalSpec(item["id"], item["assertion_id"], item["evaluator"],
                           item.get("severity", "required"), item.get("parameters", {}))
                  for item in value["evals"])
    contract = Contract(value["id"], int(value["version"]), value["intent"], (assertion,),
                        tuple(value["constraints"]), tuple(value["invariants"]), evals, scopes)
    ContractValidator().validate(contract)
    return contract


def _freeze_mapping(value):
    if isinstance(value, dict):
        return FrozenDict({key: _freeze_mapping(item) for key, item in value.items()})
    if isinstance(value, list):
        return tuple(_freeze_mapping(item) for item in value)
    return value


def _policy_evals(assertion_id, proposed, semantic_purpose, *, visual):
    output = proposed["output"]
    evals = [
        EvalSpec("artifact-valid", assertion_id, "artifact-valid"),
        EvalSpec("preservation", assertion_id, "preservation"),
        EvalSpec("nonblank", assertion_id, "nonblank"),
        EvalSpec("type", assertion_id, "type", parameters={"expected_type": output["type"]}),
        EvalSpec("output-shape", assertion_id, "output-shape", parameters={"expected_shape": output["shape"]}),
        EvalSpec("formula-errors", assertion_id, "formula-errors"),
        EvalSpec("semantic", assertion_id, "semantic-independent", parameters={
            "property": proposed["property"], "purpose": semantic_purpose,
            "uncertainty_allowed": output["uncertainty_allowed"],
        }),
    ]
    if visual:
        evals.append(EvalSpec("visual", assertion_id, "render"))
    return tuple(evals)


def _require_exact_fields(value, expected, path):
    if not isinstance(value, dict) or set(value) != expected:
        raise ValueError(f"{path} must contain exactly {sorted(expected)}")


def _state_text(value, path):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{path} must be nonempty text")
    if value.lower().lstrip().startswith(PROCEDURAL_PREFIXES):
        raise ValueError(f"procedural contract proposal rejected at {path}")
    return value.strip()


def _validate_contract_proposal(value, task, selectors):
    _require_exact_fields(value, CONTRACT_PROPOSAL_FIELDS, "contract")
    if value["schema_version"] != CONTRACT_SCHEMA_VERSION:
        raise ValueError("unsupported contract schema version")
    _require_exact_fields(value["desired_state"], {"assertions"}, "desired_state")
    assertions = value["desired_state"]["assertions"]
    if not isinstance(assertions, list) or len(assertions) != 1:
        raise ValueError("desired_state.assertions must contain exactly one assertion")
    assertion = assertions[0]
    _require_exact_fields(assertion, ASSERTION_FIELDS, "assertion")
    assertion["id"] = _state_text(assertion["id"], "assertion.id")
    assertion["description"] = _state_text(assertion["description"], "assertion.description")
    if assertion["property"] not in STATE_PROPERTIES:
        raise ValueError("assertion.property is not a supported state property")
    _require_exact_fields(assertion["output"], OUTPUT_FIELDS, "assertion.output")
    output = assertion["output"]
    if output["type"] not in OUTPUT_TYPES or output["shape"] not in OUTPUT_SHAPES:
        raise ValueError("assertion output type or shape is invalid")
    if not isinstance(output["uncertainty_allowed"], bool):
        raise ValueError("uncertainty_allowed must be boolean")
    if output["type"] == "any" and not output["uncertainty_allowed"]:
        raise ValueError("an unspecified output type must explicitly allow uncertainty")
    expected_type = task.context.get("expected_type")
    if expected_type and output["type"] not in (expected_type, "any"):
        raise ValueError("proposal output type conflicts with observed task metadata")
    expected_shape = "scalar" if _selector_cell_count(selectors) == 1 else "range"
    if output["shape"] != expected_shape:
        raise ValueError("proposal output shape conflicts with authorized selectors")
    for field in ("constraints", "invariants"):
        if not isinstance(value[field], list) or not value[field]:
            raise ValueError(f"{field} must be a nonempty list")
        value[field] = [_state_text(item, f"{field}[]") for item in value[field]]

    intents = value["evaluator_intents"]
    if not isinstance(intents, list) or not intents:
        raise ValueError("evaluator_intents must be nonempty")
    seen_evaluators = set()
    for item in intents:
        _require_exact_fields(item, EVALUATOR_INTENT_FIELDS, "evaluator_intent")
        if item["assertion_id"] != assertion["id"] or item["evaluator"] not in AVAILABLE_EVALUATORS:
            raise ValueError("evaluator intent references unavailable evaluator or assertion")
        _state_text(item["id"], "evaluator_intent.id")
        item["purpose"] = _state_text(item["purpose"], "evaluator_intent.purpose")
        if item["evaluator"] in seen_evaluators:
            raise ValueError("duplicate evaluator intent")
        seen_evaluators.add(item["evaluator"])
    required_evaluators = {"semantic-independent"}
    if task.context.get("visual_intent"):
        required_evaluators.add("render")
    if not required_evaluators <= seen_evaluators:
        raise ValueError("proposal lacks required evaluator intent")

    if any(not isinstance(item, Mapping) or not item.get("name") or not item.get("version")
           for item in task.capability_manifests):
        raise ValueError("task capability manifests must contain versioned name entries")
    available = {(str(item["name"]), str(item["version"])) for item in task.capability_manifests}
    required = value["required_capabilities"]
    if not isinstance(required, list) or not required:
        raise ValueError("required_capabilities must be nonempty")
    requested = set()
    for item in required:
        _require_exact_fields(item, CAPABILITY_REF_FIELDS, "required_capability")
        pair = (_state_text(item["name"], "capability.name"), _state_text(item["version"], "capability.version"))
        if pair not in available:
            raise ValueError(f"required capability unavailable: {pair}")
        requested.add(pair)
    if not requested & {("write_cells", VERSION), ("copy_or_fill_formula", VERSION)}:
        raise ValueError("proposal lacks a bounded mutation capability")
    return value


def _selector_cell_count(selectors):
    count = 0
    for selector in selectors:
        _, coordinates = parse_selector(selector)
        min_col, min_row, max_col, max_row = range_boundaries(coordinates)
        count += (max_col - min_col + 1) * (max_row - min_row + 1)
    return count


def _validate_transition(value, task):
    _require_exact_fields(value, TRANSITION_FIELDS, "transition")
    _require_exact_fields(value["capability"], TRANSITION_CAPABILITY_FIELDS, "transition.capability")
    name, version = value["capability"]["name"], value["capability"]["version"]
    available = {(str(item.get("name")), str(item.get("version"))) for item in task.capability_manifests
                 if isinstance(item, Mapping)}
    if (name, version) not in available:
        raise ValueError(f"transition capability unavailable: {(name, version)}")
    if name not in {"write_cells", "copy_or_fill_formula", "inspect_workbook", "recalculate",
                    "validate_workbook", "render_range"}:
        raise ValueError("transition capability is not permitted")
    if not isinstance(value["rationale"], str) or not value["rationale"].strip():
        raise ValueError("transition.rationale must be nonempty text")
    value["rationale"] = value["rationale"].strip()
    inputs = value["inputs"]
    if not isinstance(inputs, dict):
        raise ValueError("transition.inputs must be an object")
    if name == "write_cells":
        _require_exact_fields(inputs, {"writes"}, "transition.inputs")
        writes = inputs["writes"]
        if not isinstance(writes, list) or not writes or len(writes) > MAX_LITERAL_WRITES:
            raise ValueError(f"writes must contain 1..{MAX_LITERAL_WRITES} literal cells")
        for item in writes:
            _require_exact_fields(item, {"selector", "value"}, "write")
            _, coordinates = parse_selector(item["selector"])
            if ":" in coordinates:
                raise ValueError("literal write selectors must identify one cell")
    elif name == "copy_or_fill_formula":
        _require_exact_fields(inputs, {"source", "target"}, "transition.inputs")
        parse_selector(inputs["source"]); parse_selector(inputs["target"])
        if ":" in parse_selector(inputs["source"])[1]:
            raise ValueError("formula source must identify one cell")
    elif name == "inspect_workbook":
        _require_exact_fields(inputs, {"selectors"}, "transition.inputs")
        selectors = inputs["selectors"]
        if not isinstance(selectors, list) or not selectors or _selector_cell_count(selectors) > MAX_INSPECT_CELLS:
            raise ValueError("inspection transition exceeds bounded selector limit")
    elif name in {"recalculate", "validate_workbook"}:
        _require_exact_fields(inputs, set(), "transition.inputs")
    else:
        _require_exact_fields(inputs, {"selector"}, "transition.inputs")
        parse_selector(inputs["selector"])
    return value


def _json_reply(text):
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < start: raise ValueError("model reply has no JSON object")
    value = json.loads(text[start:end + 1])
    if not isinstance(value, dict): raise ValueError("model reply must be object")
    return value


def _strict_contract_reply(text):
    def object_without_duplicates(pairs):
        value = {}
        for key, item in pairs:
            if key in value:
                raise ValueError(f"duplicate contract field: {key}")
            value[key] = item
        return value
    stripped = text.strip()
    fenced = re.fullmatch(r"```(?:json)?\s*\n?(\{.*\})\s*```", stripped,
                          flags=re.IGNORECASE | re.DOTALL)
    candidate = fenced.group(1) if fenced else stripped
    try:
        value = json.loads(candidate, object_pairs_hook=object_without_duplicates)
    except json.JSONDecodeError as error:
        raise ValueError("contract proposal must be exactly one JSON object") from error
    if not isinstance(value, dict):
        raise ValueError("contract proposal must be a JSON object")
    return value


def _strict_semantic_verdict(text):
    value = _json_reply(text)
    required = {"verdict", "expected_state", "rationale", "confidence"}
    if set(value) != required:
        raise ValueError("semantic verdict must contain exactly the required fields")
    if value["verdict"] not in {"pass", "fail", "uncertain"}:
        raise ValueError("semantic verdict is invalid")
    if not isinstance(value["rationale"], str) or not value["rationale"].strip():
        raise ValueError("semantic rationale must be nonempty")
    confidence = value["confidence"]
    if isinstance(confidence, bool) or not isinstance(confidence, (int, float)) or not 0 <= confidence <= 1:
        raise ValueError("semantic confidence must be between zero and one")
    return value


def _model_context(context):
    """Keep independent evaluator inputs isolated from contract/action generation."""
    return {key: value for key, value in context.items() if key != "independent_expected"}


def _selected_cells(path, selectors):
    wb = load_workbook(path, data_only=False); rows = []
    for selector in selectors:
        sheet, coordinates = parse_selector(selector)
        min_col, min_row, max_col, max_row = range_boundaries(coordinates)
        for row in wb[sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row: rows.append({"selector": f"{sheet}!{cell.coordinate}", "value": _cell_value(cell.value)})
    wb.close(); return rows


def _all_cells(path):
    wb = load_workbook(path, data_only=False); result = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None: result[f"workbook/{ws.title}/{cell.coordinate}"] = _cell_value(cell.value)
    wb.close(); return result


def _valid(path):
    try:
        wb = load_workbook(path); wb.close(); return True, None
    except Exception as error: return False, f"{type(error).__name__}:{error}"


def _cell_value(value):
    """Keep spreadsheet implementation objects below the generic state boundary."""
    if isinstance(value, ArrayFormula):
        return value.text
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _type_name(value):
    if isinstance(value, bool): return "boolean"
    if isinstance(value, (int, float)): return "number"
    if isinstance(value, str) and value.startswith("="): return "formula"
    if isinstance(value, str): return "text"
    if value is None: return "blank"
    if hasattr(value, "isoformat") and hasattr(value, "year"): return "date"
    return "other"


def _file_hash(path):
    import hashlib
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _eval_mapping(result):
    return {"eval_id": result.eval_id, "status": result.status.value, "message": result.message,
            "details": dict(result.details)}


def _discrepancy_mapping(item):
    return {"id": item.id, "kind": item.kind.value, "failed_eval": item.failed_eval_id,
            "expected": item.expected, "observed": item.observed,
            "permissible_mutation_scope": [s.resource for s in item.permissible_mutation_scope]}
