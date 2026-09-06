"""SpreadsheetBench factory with bounded context and no fulfilment-time goldens."""

from __future__ import annotations

import json
import os
import urllib.error
import urllib.request
from dataclasses import asdict
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from adapters.spreadsheet import KIND, manifests
from experiment.runner import ModelReply, Task
from services.spreadsheet import SpreadsheetServices


class ProviderConfigurationError(RuntimeError):
    pass


class ProviderRequestError(RuntimeError):
    pass


class OpenRouterProvider:
    """Small synchronous OpenAI-compatible provider used by TaskRuntime."""

    endpoint = "https://openrouter.ai/api/v1/chat/completions"
    provider_name = "OpenRouter"

    def __init__(self, api_key: str | None = None, *, timeout_seconds: int = 120):
        self.api_key = api_key or os.environ.get("OPENROUTER_API_KEY")
        self.timeout_seconds = timeout_seconds

    def complete(self, prompt: str, *, model: str, temperature: float) -> ModelReply:
        if not self.api_key:
            raise ProviderConfigurationError("OPENROUTER_API_KEY is required")
        body = json.dumps({"model": model, "temperature": temperature,
                           "messages": [{"role": "user", "content": prompt}]}).encode()
        request = urllib.request.Request(self.endpoint, data=body, method="POST", headers={
            "Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json",
            "HTTP-Referer": "https://github.com/ylookup/encode-hackathon",
            "X-Title": "Declarative fulfilment experiment",
        })
        try:
            with urllib.request.urlopen(request, timeout=self.timeout_seconds) as response:
                payload = json.loads(response.read())
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as error:
            raise ProviderRequestError(f"{self.provider_name} request failed: {error}") from error
        try:
            text = payload["choices"][0]["message"]["content"]
        except (KeyError, IndexError, TypeError) as error:
            raise ProviderRequestError(f"{self.provider_name} response lacks message content") from error
        usage = payload.get("usage") or {}
        return ModelReply(str(text), int(usage.get("prompt_tokens") or 0),
                          int(usage.get("completion_tokens") or 0), float(usage.get("cost") or 0),
                          payload.get("id"))


class GeminiProvider(OpenRouterProvider):
    """Direct Gemini API through Google's OpenAI-compatible endpoint."""

    endpoint = "https://generativelanguage.googleapis.com/v1beta/openai/chat/completions"
    provider_name = "Gemini"

    def __init__(self, api_key: str | None = None, *, timeout_seconds: int = 120):
        self.api_key = api_key or os.environ.get("GEMINI_API_KEY")
        self.timeout_seconds = timeout_seconds

    def complete(self, prompt: str, *, model: str, temperature: float) -> ModelReply:
        if not self.api_key:
            raise ProviderConfigurationError("GEMINI_API_KEY is required")
        return super().complete(prompt, model=model, temperature=temperature)


class ScriptedProvider:
    """Deterministic CLI smoke-test provider; outputs are not benchmark evidence."""

    def __init__(self, replies: list[str]):
        self.replies = list(replies)
        self.prompts: list[str] = []

    def complete(self, prompt: str, *, model: str, temperature: float) -> ModelReply:
        self.prompts.append(prompt)
        if not self.replies:
            raise ProviderConfigurationError("scripted provider has no reply remaining")
        return ModelReply(self.replies.pop(0), 0, 0, 0.0, "scripted-not-benchmark-evidence")


def bounded_workbook_context(path: Path, *, max_cells: int = 400, max_chars: int = 30_000) -> dict[str, Any]:
    """Canonical, explicitly truncated facts; loads formulas, never cached goldens."""
    if max_cells < 1 or max_chars < 100:
        raise ValueError("context limits are too small")
    wb = load_workbook(path, data_only=False, read_only=False)
    cells = []
    omitted_nonempty = 0
    chars = 0
    stopped = False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                record = {"selector": f"{ws.title}!{cell.coordinate}", "value": _json_value(cell.value),
                          "data_type": cell.data_type, "style": {"style_id": cell.style_id,
                          "number_format": cell.number_format}}
                encoded = json.dumps(record, sort_keys=True, ensure_ascii=False)
                if len(cells) >= max_cells or chars + len(encoded) > max_chars:
                    omitted_nonempty += 1; stopped = True
                    continue
                cells.append(record); chars += len(encoded)
    sheet_names = list(wb.sheetnames); wb.close()
    return {"schema_version": "1.0.0", "artifact_kind": KIND, "facts": {"sheet_names": sheet_names,
            "cells": cells}, "interpretations": {}, "inspected_scope": ["workbook/*"],
            "omitted_scope": ["workbook/unserialized/*"] if stopped else [],
            "truncation": {"truncated": stopped, "max_cells": max_cells, "max_chars": max_chars,
                           "serialized_cells": len(cells), "omitted_nonempty_cells": omitted_nonempty,
                           "serialized_chars": chars}}


def factory(raw: Mapping[str, Any]):
    """CLI factory returning services, provider, and runtime-safe task records."""
    backend = dict(raw.get("backend_config") or {})
    if float(raw.get("run_config", {}).get("temperature", 0)) != 0:
        raise ProviderConfigurationError("preregistered backend requires temperature 0")
    dataset_dir = Path(backend.get("dataset_dir") or raw.get("dataset_dir") or "/data")
    embedded_selection = raw.get("task_selection")
    selection = backend.get("selection_manifest") or raw.get("selection_manifest")
    selected_ids = (_selection_ids_value(embedded_selection) if embedded_selection is not None
                    else _selection_ids(Path(selection)) if selection else None)
    limits = backend.get("context_limits") or {}
    tasks = load_runtime_tasks(dataset_dir, selected_ids=selected_ids,
                               max_cells=int(limits.get("max_cells", 400)),
                               max_chars=int(limits.get("max_chars", 30_000)))
    provider_config = backend.get("provider") or {"type": "openrouter"}
    provider_type = provider_config.get("type", "openrouter")
    if provider_type == "openrouter":
        provider = OpenRouterProvider(timeout_seconds=int(provider_config.get("timeout_seconds", 120)))
    elif provider_type == "gemini":
        provider = GeminiProvider(timeout_seconds=int(provider_config.get("timeout_seconds", 120)))
    elif provider_type == "scripted":
        provider = ScriptedProvider(list(provider_config.get("replies") or []))
    else:
        raise ProviderConfigurationError(f"unknown provider type: {provider_type}")
    return SpreadsheetServices(), provider, tasks


def load_runtime_tasks(dataset_dir: Path, *, selected_ids: set[str] | None,
                       max_cells: int, max_chars: int) -> list[Task]:
    records = json.loads((dataset_dir / "dataset.json").read_text())
    manifest_by_name = tuple(_manifest_mapping(item) for item in manifests())
    tasks = []
    for record in records:
        task_id = str(record["id"])
        if selected_ids is not None and task_id not in selected_ids:
            continue
        folder = dataset_dir / record["spreadsheet_path"]
        candidates = sorted(folder.glob("*init*.xlsx"))
        if len(candidates) != 1:
            raise FileNotFoundError(f"task {task_id} must have exactly one init workbook")
        init_path = candidates[0]
        context = {key: record[key] for key in ("answer_sheet", "answer_position", "instruction_type", "data_position") if key in record}
        context["workbook_observation"] = bounded_workbook_context(init_path, max_cells=max_cells, max_chars=max_chars)
        tasks.append(Task(task_id, record["instruction"], init_path, KIND, context, manifest_by_name))
    if selected_ids is not None and {task.id for task in tasks} != selected_ids:
        missing = selected_ids - {task.id for task in tasks}
        raise ValueError(f"selection contains unknown task ids: {sorted(missing)}")
    return tasks


def _selection_ids(path: Path) -> set[str]:
    return _selection_ids_value(json.loads(path.read_text()))


def _selection_ids_value(value) -> set[str]:
    rows = value.get("task_ids") if isinstance(value, dict) else value
    if rows is None and isinstance(value, dict) and isinstance(value.get("tasks"), list):
        rows = [row.get("id") if isinstance(row, dict) else row for row in value["tasks"]]
    if not isinstance(rows, list) or not rows:
        raise ValueError("selection manifest must contain a non-empty task_ids list")
    ids = {str(item) for item in rows}
    if len(ids) != len(rows): raise ValueError("selection task ids must be unique")
    return ids


def _manifest_mapping(manifest):
    value = asdict(manifest)
    value["effect"] = manifest.effect.value
    value["possible_mutation_scope"] = [scope.resource for scope in manifest.possible_mutation_scope]
    return value


def _json_value(value):
    if isinstance(value, ArrayFormula): return value.text
    if value is None or isinstance(value, (bool, int, float, str)): return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)
