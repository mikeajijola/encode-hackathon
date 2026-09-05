import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from adapters.spreadsheet import KIND
from experiment.runner import Arm, ExperimentRunner, ModelReply, RunConfig, Task
from services.spreadsheet import MAX_LITERAL_WRITES, SpreadsheetServices
from sb import answer_cells
from tests.contract_fixtures import capability_records, contract_reply, transition_reply


class Provider:
    def __init__(self, replies): self.replies, self.prompts = list(replies), []
    def complete(self, prompt, *, model, temperature):
        self.prompts.append(prompt)
        return ModelReply(self.replies.pop(0), 1, 1, 0, f"p-{len(self.prompts)}")


def config(arm):
    return RunConfig("capabilities", arm, "same", "v1", 0, 1000, 8, 10000, 1,
                     "sha256:x", "none", "fixed", {"transport": 0})


class CapabilityPlanningTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.path = self.root / "initial.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        for row in range(1, 102): ws.cell(row, 1).value = row
        ws["B1"] = "=A1*2"; ws["D1"] = "preserve"
        wb.save(self.path); wb.close()

    def tearDown(self): self.temp.cleanup()

    def task(self, answer="B2:B100"):
        context = {"answer_sheet": "Data", "answer_position": answer, "expected_type": "formula"}
        return Task("range", "fill a relative doubled formula down the answer range",
                    self.path, KIND, context, capability_records())

    def run_case(self, arm, replies, *, answer="B2:B100", service=None):
        provider = Provider(replies); out = self.root / f"out-{len(list(self.root.glob('out-*')))}"
        result = ExperimentRunner(config(arm), service or SpreadsheetServices(), provider, out).run(
            [self.task(answer)])[0]
        return out, provider, result

    def formula_contract(self, shape="range"):
        return contract_reply("every target cell contains the corresponding relative doubled formula",
                              property="formula_result", output_type="formula", shape=shape,
                              capabilities=["copy_or_fill_formula"])

    def test_compact_formula_fill_expands_to_exact_authorized_cells(self):
        transition = transition_reply("copy_or_fill_formula",
            {"source": "Data!B1", "target": "Data!B2:B100"})
        out, provider, result = self.run_case(Arm.B, [self.formula_contract(), transition])
        self.assertEqual(result["status"], "ok")
        self.assertLess(len(transition), 250)
        wb = load_workbook(out / result["output"], data_only=False)
        self.assertEqual(wb["Data"]["B2"].value, "=A2*2")
        self.assertEqual(wb["Data"]["B100"].value, "=A100*2")
        self.assertEqual(wb["Data"]["D1"].value, "preserve"); wb.close()
        events = [json.loads(line) for line in (out / "events" / "range.jsonl").read_text().splitlines()]
        capability = next(event for event in events if event["event_type"] == "capability_result")
        self.assertEqual(capability["payload"]["capability"], "copy_or_fill_formula")
        self.assertEqual(len(capability["payload"]["actual_scope"]), 99)
        self.assertEqual(capability["payload"]["provenance"]["adapter_version"], "1.0.0")

    def test_compact_range_cannot_expand_beyond_contract_scope(self):
        transition = transition_reply("copy_or_fill_formula",
            {"source": "Data!B1", "target": "Data!B2:B101"})
        out, _, result = self.run_case(Arm.B, [self.formula_contract(), transition])
        self.assertEqual(result["status"], "action_failed")
        wb = load_workbook(out / result["output"])
        self.assertIsNone(wb["Data"]["B2"].value); self.assertIsNone(wb["Data"]["B101"].value); wb.close()
        self.assertIn("scope_violation_rejected", (out / "events" / "range.jsonl").read_text())

    def test_arbitrary_execute_capability_is_rejected(self):
        transition = transition_reply("execute", {"code": "arbitrary"})
        _, _, result = self.run_case(Arm.A, [transition])
        self.assertTrue(result["status"].startswith("error: ValueError: transition capability unavailable"))

    def test_known_capability_missing_from_task_manifest_is_rejected(self):
        records = tuple(record for record in capability_records()
                        if record["name"] != "copy_or_fill_formula")
        task = Task("missing-capability", "fill formulas", self.path, KIND,
                    {"answer_sheet": "Data", "answer_position": "B2:B100",
                     "expected_type": "formula"}, records)
        provider = Provider([transition_reply("copy_or_fill_formula",
                            {"source": "Data!B1", "target": "Data!B2:B100"})])
        result = ExperimentRunner(config(Arm.A), SpreadsheetServices(), provider,
                                  self.root / "missing-out").run([task])[0]
        self.assertTrue(result["status"].startswith(
            "error: ValueError: transition capability unavailable"))

    def test_recalculation_and_validation_failure_cannot_succeed(self):
        recalc = transition_reply("recalculate", {})
        with patch("adapters.spreadsheet.shutil.which", return_value=None):
            _, _, result = self.run_case(Arm.A, [recalc])
        self.assertEqual(result["status"], "action_failed")

        wb = load_workbook(self.path); wb["Data"]["C1"] = "#VALUE!"; wb.save(self.path); wb.close()
        validate = transition_reply("validate_workbook", {})
        _, _, result = self.run_case(Arm.A, [validate])
        self.assertEqual(result["status"], "action_failed")

    def test_inspection_is_bounded_and_same_manifests_reach_arm_a(self):
        inspect = transition_reply("inspect_workbook", {"selectors": ["Data!A1:A20"]})
        out, provider, result = self.run_case(Arm.A, [inspect])
        self.assertEqual(result["status"], "ok")
        prompt = json.loads(provider.prompts[0])
        self.assertEqual(prompt["capability_manifests"], json.loads(json.dumps(capability_records(), default=str)))
        capability = next(json.loads(line) for line in (out / "events" / "range.jsonl").read_text().splitlines()
                          if json.loads(line)["event_type"] == "capability_result")
        self.assertIn("observation", capability["payload"]["output"])

    def test_inspection_observation_feeds_next_reconciliation_transition(self):
        context = {"answer_sheet": "Data", "answer_position": "B2", "expected_type": "number",
                   "independent_expected": 4}
        task = Task("inspect-loop", "derive twice A2 into B2", self.path, KIND, context, capability_records())
        provider = Provider([
            contract_reply("B2 equals twice A2", output_type="number"),
            transition_reply("inspect_workbook", {"selectors": ["Data!A2"]}),
            transition_reply(inputs={"writes": [{"selector": "Data!B2", "value": 4}]}),
        ])
        out = self.root / "inspect-loop-out"
        result = ExperimentRunner(config(Arm.D), SpreadsheetServices(), provider, out).run([task])[0]
        self.assertEqual(result["status"], "fulfilled")
        second_plan = json.loads(provider.prompts[2])
        self.assertEqual(second_plan["capability_observations"][0]["capability"], "inspect_workbook")
        cells = second_plan["capability_observations"][0]["output"]["observation"]["facts"]["cells"]
        self.assertEqual(cells[0]["value"], 2)

    def test_literal_write_limit_prevents_large_action_payload(self):
        writes = [{"selector": f"Data!B{row}", "value": row} for row in range(2, MAX_LITERAL_WRITES + 3)]
        transition = transition_reply(inputs={"writes": writes})
        _, _, result = self.run_case(Arm.A, [transition])
        self.assertTrue(result["status"].startswith("error: ValueError: writes must contain"))

    def test_committed_development_large_target_audit_is_reproducible(self):
        protocol = Path(__file__).resolve().parents[1] / "protocol"
        selection = json.loads((protocol / "development_selection.json").read_text())
        audit = json.loads((protocol / "development_target_audit.json").read_text())
        sizes = {task["id"]: len(answer_cells(task)) for task in selection["tasks"]}
        self.assertEqual(audit["selection_sha256"], selection["selection_sha256"])
        self.assertEqual(audit["tasks"], len(sizes))
        self.assertEqual(audit["maximum_target_cells"], max(sizes.values()))
        self.assertEqual((audit["maximum_task_id"], sizes[audit["maximum_task_id"]]), ("209-30", 6066))
        expected_large = sorted((task_id, size) for task_id, size in sizes.items() if size > MAX_LITERAL_WRITES)
        observed_large = sorted((row["id"], row["target_cells"]) for row in audit["targets_over_literal_write_limit"])
        self.assertEqual(observed_large, expected_large)


if __name__ == "__main__":
    unittest.main()
