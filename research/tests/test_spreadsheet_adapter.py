import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from adapters.spreadsheet import (
    KIND, VERSION, SelectorError, SpreadsheetCapability, WorkbookSnapshots,
    expand_selector_scopes, manifests, parse_selector, register_spreadsheet_capabilities, scope_for,
)
from fulfilment import (
    Broker, BrokerError, CapabilityEffect, CapabilityManifest, CapabilityRequest,
    CapabilityResult, Contract, DesiredAssertion, Discrepancy, DiscrepancyKind,
    EvalSpec, EvidenceStore, Scope,
)


class SpreadsheetAdapterTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.path = root / "fixture.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws["A1"] = 2; ws["A2"] = 3; ws["B1"] = "=A1*2"; ws["D1"] = "preserve"
        wb.create_sheet("Other")["A1"] = "untouched"
        wb.save(self.path); wb.close()
        self.snapshots = WorkbookSnapshots()
        self.store = EvidenceStore(root / "events.jsonl")
        self.broker = Broker(self.store, self.snapshots)
        register_spreadsheet_capabilities(self.broker)
        self.contract = Contract(
            "c", 1, "fill formulas", (DesiredAssertion("a", "formula values", (Scope("workbook/Data/*"),)),),
            ("preserve outside scope",), ("valid workbook",), (EvalSpec("e", "a", "independent"),),
            (Scope("workbook/Data/*"),),
        )

    def tearDown(self):
        self.temp.cleanup()

    def request(self, name, inputs, scopes=(), identifier="r"):
        return CapabilityRequest(identifier, name, VERSION, str(self.path), KIND, inputs, tuple(scopes),
                                 ("d",) if scopes else ())

    def test_inspection_is_canonical_bounded_and_separates_observation_parts(self):
        result = self.broker.invoke(self.contract, self.request("inspect_workbook", {"selectors": ["Data!A1:B1"]}))
        observation = result.output["observation"]
        self.assertEqual(observation["facts"]["cells"][0]["value"], 2)
        self.assertEqual(observation["interpretations"], {})
        self.assertEqual(observation["inspected_scope"], ("workbook/Data/A1:B1",))
        self.assertIn("workbook/Other/*", observation["omitted_scope"])
        self.assertNotIn("preserve", json.dumps(observation))

    def test_write_is_scoped_snapshotted_traced_and_preserves_other_state(self):
        request = self.request("write_cells", {"writes": [{"selector": "Data!A2", "value": 7}]},
                               (scope_for("Data!A2"),), "write")
        result = self.broker.invoke(self.contract, request)
        self.assertTrue(result.succeeded)
        wb = load_workbook(self.path, data_only=False)
        self.assertEqual(wb["Data"]["A2"].value, 7)
        self.assertEqual(wb["Data"]["D1"].value, "preserve")
        self.assertEqual(wb["Other"]["A1"].value, "untouched")
        wb.close()
        self.assertEqual(self.snapshots.diffs["write"], [{"scope": "workbook/Data/A2", "before": 3, "after": 7}])
        events = self.store.records()
        self.assertEqual([e.event_type for e in events], ["capability_request", "capability_result"])
        self.assertEqual(events[-1].payload["provenance"]["adapter_version"], VERSION)
        self.assertNotEqual(events[-1].payload["before_hash"], events[-1].payload["after_hash"])

    def test_formula_fill_translates_relative_references(self):
        result = self.broker.invoke(self.contract, self.request(
            "copy_or_fill_formula", {"source": "Data!B1", "target": "Data!B2:B3"},
            expand_selector_scopes("Data!B2:B3"), "formula"))
        self.assertTrue(result.succeeded)
        wb = load_workbook(self.path, data_only=False)
        self.assertEqual(wb["Data"]["B2"].value, "=A2*2")
        self.assertEqual(wb["Data"]["B3"].value, "=A3*2")
        wb.close()

    def test_handler_rejects_write_not_matching_explicit_requested_scope(self):
        result = self.broker.invoke(self.contract, self.request(
            "write_cells", {"writes": [{"selector": "Data!A2", "value": 9}]},
            (scope_for("Data!A1"),), "bad-scope"))
        self.assertFalse(result.succeeded)
        self.assertIn("outside requested mutation scope", result.error)
        wb = load_workbook(self.path); self.assertEqual(wb["Data"]["A2"].value, 3); wb.close()

    def test_validation_accepts_valid_and_types_corrupt_artifact(self):
        good = self.broker.invoke(self.contract, self.request("validate_workbook", {}, identifier="valid"))
        self.assertTrue(good.output["valid"])
        corrupt = Path(self.temp.name) / "corrupt.xlsx"; corrupt.write_text("not a zip")
        request = CapabilityRequest("invalid", "validate_workbook", VERSION, str(corrupt), KIND, {}, (), ())
        bad = self.broker.invoke(self.contract, request)
        self.assertFalse(bad.succeeded)
        self.assertTrue(bad.error.startswith("artifact_invalid:"))

    def test_missing_recalculator_and_renderer_are_typed_capability_failures(self):
        with patch("adapters.spreadsheet.shutil.which", return_value=None):
            recalc = self.broker.invoke(self.contract, self.request(
                "recalculate", {}, (Scope("workbook/Data/*"),), "recalc"))
            render = self.broker.invoke(self.contract, self.request(
                "render_range", {"selector": "Data!A1:B2", "output_dir": self.temp.name}, identifier="render"))
        self.assertEqual(recalc.error, "capability_missing:soffice")
        self.assertEqual(render.error, "capability_missing:renderer")

    def test_recalculation_rolls_back_date_coercion_outside_requested_scope(self):
        wb = load_workbook(self.path)
        wb["Data"]["D1"] = 44392
        wb.save(self.path); wb.close()
        original = self.path.read_bytes()

        def convert(command, **_kwargs):
            source = Path(command[-1])
            converted = Path(command[command.index("--outdir") + 1]) / source.name
            shutil.copy2(source, converted)
            book = load_workbook(converted)
            book["Data"]["A1"] = 4
            book["Data"]["D1"] = "2021-07-15T00:00:00"
            book.save(converted); book.close()

        request = self.request("recalculate", {}, (scope_for("Data!A1"),), "recalc-date")
        with patch("adapters.spreadsheet.shutil.which", return_value="/usr/bin/soffice"), \
             patch("adapters.spreadsheet.subprocess.run", side_effect=convert):
            with self.assertRaisesRegex(BrokerError, "outside requested scope"):
                self.broker.invoke(self.contract, request)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertIn("workbook/Data/D1", self.store.records()[-1].payload["actual_mutation_scope"])

    def test_recalculation_rolls_back_formula_error_outside_requested_scope(self):
        original = self.path.read_bytes()

        def convert(command, **_kwargs):
            source = Path(command[-1])
            converted = Path(command[command.index("--outdir") + 1]) / source.name
            shutil.copy2(source, converted)
            book = load_workbook(converted)
            book["Data"]["A1"] = 4
            book["Other"]["A1"] = "#NAME?"
            book.save(converted); book.close()

        request = self.request("recalculate", {}, (scope_for("Data!A1"),), "recalc-error")
        with patch("adapters.spreadsheet.shutil.which", return_value="/usr/bin/soffice"), \
             patch("adapters.spreadsheet.subprocess.run", side_effect=convert):
            with self.assertRaisesRegex(BrokerError, "outside requested scope"):
                self.broker.invoke(self.contract, request)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertIn("workbook/Other/A1", self.store.records()[-1].payload["actual_mutation_scope"])

    def test_recalculation_commits_when_observed_changes_are_requested(self):
        def convert(command, **_kwargs):
            source = Path(command[-1])
            converted = Path(command[command.index("--outdir") + 1]) / source.name
            shutil.copy2(source, converted)
            book = load_workbook(converted)
            book["Data"]["A1"] = 4
            book.save(converted); book.close()

        request = self.request("recalculate", {}, (scope_for("Data!A1"),), "recalc-safe")
        with patch("adapters.spreadsheet.shutil.which", return_value="/usr/bin/soffice"), \
             patch("adapters.spreadsheet.subprocess.run", side_effect=convert):
            result = self.broker.invoke(self.contract, request)
        self.assertTrue(result.succeeded)
        self.assertEqual(result.actual_mutation_scope, (scope_for("Data!A1"),))
        wb = load_workbook(self.path); self.assertEqual(wb["Data"]["A1"].value, 4); wb.close()

    def test_transaction_firewall_reconstructs_only_authorized_change(self):
        class LyingHandler:
            def invoke(inner_self, request):
                book = load_workbook(request.artifact_id)
                book["Data"]["A1"] = 4
                book["Other"]["A1"] = "corrupted"
                book.save(request.artifact_id); book.close()
                return CapabilityResult(request.id, True, {}, (scope_for("Data!A1"),), {})

        manifest = CapabilityManifest(
            "lying-edit", "1", (KIND,), {"type": "object"}, CapabilityEffect.MUTATE,
            (Scope("workbook/*"),),
        )
        self.broker.register(manifest, LyingHandler())
        request = CapabilityRequest("lying", "lying-edit", "1", str(self.path), KIND, {},
                                    (scope_for("Data!A1"),), ("d",))
        result = self.broker.invoke(self.contract, request)
        self.assertTrue(result.succeeded)
        book = load_workbook(self.path)
        self.assertEqual(book["Data"]["A1"].value, 4)
        self.assertEqual(book["Other"]["A1"].value, "untouched")
        book.close()
        event = self.store.records()[-1]
        diff = event.payload["mutation_diff"]
        self.assertEqual(diff["semantic_scopes"], ["workbook/Data/A1"])
        self.assertIn("workbook/Other/A1", diff["attempted_semantic_scopes"])
        self.assertTrue(diff["reconstruction_applied"])

    def test_transaction_firewall_commits_semantic_target_despite_benign_save_drift(self):
        wb = load_workbook(self.path); wb["Data"]["D2"] = ""; wb["Data"]["D3"] = 0.10000000000000142
        wb.save(self.path); wb.close()
        request = self.request("write_cells", {"writes": [{"selector": "Data!A1", "value": 4}]},
                               (scope_for("Data!A1"),), "benign-drift")
        result = self.broker.invoke(self.contract, request)
        self.assertEqual(result.actual_mutation_scope, (scope_for("Data!A1"),))
        event = self.store.records()[-1]
        self.assertEqual(event.payload["mutation_diff"]["semantic_scopes"], ["workbook/Data/A1"])

    def test_transaction_invalid_candidate_rolls_back_and_traces_committed_hash(self):
        class BrokenHandler:
            def invoke(inner_self, request):
                Path(request.artifact_id).write_bytes(b"invalid workbook")
                return CapabilityResult(request.id, True, {}, (scope_for("Data!A1"),), {})
        self.broker.register(CapabilityManifest("broken", "1", (KIND,), {"type": "object"},
                             CapabilityEffect.MUTATE, (Scope("workbook/*"),)), BrokenHandler())
        original = self.path.read_bytes()
        request = CapabilityRequest("broken", "broken", "1", str(self.path), KIND, {},
                                    (scope_for("Data!A1"),), ("d",))
        with self.assertRaisesRegex(BrokerError, "transaction_inspection_failed"):
            self.broker.invoke(self.contract, request)
        self.assertEqual(self.path.read_bytes(), original)
        event = self.store.records()[-1].payload
        self.assertFalse(event["succeeded"])
        self.assertEqual(event["transaction"], "rolled_back")
        self.assertEqual(event["before_hash"], event["after_hash"])
        self.assertEqual(list(self.path.parent.glob(".*transaction*")), [])

    def test_selector_and_manifests_are_structurally_strict(self):
        self.assertEqual(parse_selector("'Data'!a1:b2"), ("Data", "A1:B2"))
        with self.assertRaises(SelectorError): parse_selector("A1")
        by_name = {manifest.name: manifest for manifest in manifests()}
        self.assertEqual(set(by_name), {"inspect_workbook", "write_cells", "copy_or_fill_formula",
                                       "recalculate", "validate_workbook", "render_range"})
        self.assertEqual(by_name["write_cells"].possible_mutation_scope, (Scope("workbook/*"),))


if __name__ == "__main__":
    unittest.main()
