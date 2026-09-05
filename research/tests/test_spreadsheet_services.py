import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from adapters.spreadsheet import KIND
from experiment.runner import Arm, ExperimentRunner, ModelReply, RunConfig, Task
from services.spreadsheet import SpreadsheetServices
from tests.contract_fixtures import capability_records, contract_reply, transition_reply


class QueueProvider:
    def __init__(self, replies):
        self.replies = list(replies); self.calls = []
    def complete(self, prompt, *, model, temperature):
        self.calls.append(prompt)
        return ModelReply(self.replies.pop(0), 3, 2, .001, f"fake-{len(self.calls)}")


def config(arm):
    return RunConfig("services-fixture", arm, "fake", "fake-v1", 0, 100, 6, 10_000, 1,
                     "sha256:fixture", "none", "fixed", {"transport": 0})


class SpreadsheetServicesE2E(unittest.TestCase):
    def fixture(self, root):
        path = root / "initial.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws["A1"] = 2; ws["B1"] = None; ws["D1"] = "preserve"
        wb.save(path); wb.close()
        return path

    def run_arm(self, arm, replies, context=None):
        td = tempfile.TemporaryDirectory(); root = Path(td.name)
        artifact = self.fixture(root); out = root / "out"
        task_context = {"answer_sheet": "Data", "answer_position": "B1", "expected_type": "number",
                        "independent_expected": 4}
        task_context.update(context or {})
        provider = QueueProvider(replies)
        result = ExperimentRunner(config(arm), SpreadsheetServices(), provider, out).run(
            [Task("t1", "put twice A1 in B1", artifact, KIND, task_context, capability_records())])[0]
        return td, out, provider, result

    def value(self, out, result):
        wb = load_workbook(out / result["output"], data_only=False)
        value, preserved = wb["Data"]["B1"].value, wb["Data"]["D1"].value
        wb.close(); return value, preserved

    def test_arm_c_detects_fault_once_and_d_repairs_from_discrepancy(self):
        contract = contract_reply("B1 equals independently computed twice A1")
        wrong = transition_reply(inputs={"writes": [{"selector": "Data!B1", "value": "wrong"}]})
        correct = transition_reply()

        td_c, out_c, provider_c, c = self.run_arm(Arm.C, [contract, wrong])
        self.addCleanup(td_c.cleanup)
        self.assertEqual(self.value(out_c, c), ("wrong", "preserve"))
        self.assertFalse(c["terminal_eval"]["passed"])
        self.assertEqual(len(provider_c.calls), 2)  # no repair after terminal C eval
        failed_ids = {e["eval_id"] for e in c["terminal_eval"]["details"]["evals"] if e["status"] != "pass"}
        self.assertEqual(failed_ids, {"type", "semantic"})

        td_d, out_d, provider_d, d = self.run_arm(Arm.D, [contract, wrong, correct])
        self.addCleanup(td_d.cleanup)
        self.assertEqual(self.value(out_d, d), (4, "preserve"))
        self.assertEqual(d["status"], "fulfilled")
        self.assertTrue(d["terminal_eval"]["passed"])
        self.assertEqual(len(provider_d.calls), 3)
        self.assertFalse(any('"independent_expected": 4' in prompt for prompt in provider_d.calls))
        events = [json.loads(line) for line in (out_d / "events" / "t1.jsonl").read_text().splitlines()]
        kinds = [event["event_type"] for event in events]
        self.assertIn("discrepancy", kinds); self.assertIn("capability_result", kinds)
        capability = next(e for e in events if e["event_type"] == "capability_result")
        self.assertEqual(capability["payload"]["actual_scope"], ["workbook/Data/B1"])
        broker_events = out_d / "events" / "t1.broker.jsonl"
        self.assertTrue(broker_events.exists())
        self.assertIn("artifact_sha256", broker_events.read_text())

    def test_all_arms_use_provider_interface_and_preserve_treatments(self):
        direct = transition_reply()
        contract = contract_reply("B1 has the requested state")
        cases = {Arm.A: [direct], Arm.B: [contract, direct]}
        for arm, replies in cases.items():
            with self.subTest(arm=arm):
                td, out, provider, result = self.run_arm(arm, replies)
                self.addCleanup(td.cleanup)
                self.assertEqual(self.value(out, result), (4, "preserve"))
                self.assertEqual(len(provider.calls), 1 if arm is Arm.A else 2)
                self.assertEqual(result["contract_compiled"], arm is not Arm.A)

    def test_absent_independent_semantics_is_typed_uncertain_not_fulfilled(self):
        contract = contract_reply("B1 answers the intent")
        td = tempfile.TemporaryDirectory(); self.addCleanup(td.cleanup)
        root = Path(td.name); artifact = self.fixture(root)
        context = {"answer_sheet": "Data", "answer_position": "B1", "expected_type": "number"}
        out2 = root / "out"; provider2 = QueueProvider([
            contract, transition_reply(), '{"verdict":"pass"}'
        ])
        uncertain = ExperimentRunner(config(Arm.D), SpreadsheetServices(), provider2, out2).run(
            [Task("u1", "derive an answer", artifact, KIND, context, capability_records())])[0]
        self.assertEqual(uncertain["status"], "unfulfilled:evaluation_uncertain")
        self.assertFalse(uncertain["terminal_eval"]["passed"])
        discrepancies = uncertain["terminal_eval"]["details"]["discrepancies"]
        self.assertTrue(any(d["kind"] == "evaluation_uncertainty" for d in discrepancies))
        self.assertEqual(len(provider2.calls), 3)  # contract, first action, malformed independent eval

    def test_scope_violation_is_rejected_without_mutation(self):
        contract = contract_reply("B1 is populated")
        malicious = transition_reply(inputs={"writes": [{"selector": "Data!D1", "value": "changed"}]})
        td, out, provider, result = self.run_arm(Arm.B, [contract, malicious])
        self.addCleanup(td.cleanup)
        self.assertEqual(self.value(out, result), (None, "preserve"))
        events = (out / "events" / "t1.jsonl").read_text()
        self.assertIn("scope_violation_rejected", events)

    def test_visual_intent_missing_renderer_is_nonpass(self):
        base = json.loads(contract_reply("B1 is visually and semantically correct", property="visual_state"))
        base["evaluator_intents"].append({"id": "visual-check", "assertion_id": "answer-state",
            "evaluator": "render", "purpose": "verify the requested visual state"})
        contract = json.dumps(base)
        with patch("adapters.spreadsheet.shutil.which", return_value=None):
            td, out, provider, result = self.run_arm(Arm.C,
                [contract, transition_reply()], {"visual_intent": True})
        self.addCleanup(td.cleanup)
        self.assertFalse(result["terminal_eval"]["passed"])
        visual = next(e for e in result["terminal_eval"]["details"]["evals"] if e["eval_id"] == "visual")
        self.assertEqual(visual["status"], "fail")
        self.assertIn("capability_missing:renderer", visual["message"])

    def test_golden_like_runtime_context_is_rejected(self):
        td, out, provider, result = self.run_arm(Arm.B, [], {"golden_xlsx": "/forbidden.xlsx"})
        self.addCleanup(td.cleanup)
        self.assertTrue(result["status"].startswith("error: ValueError"))
        self.assertEqual(len(provider.calls), 0)

    def test_null_answer_sheet_resolves_to_active_sheet(self):
        td, out, provider, result = self.run_arm(
            Arm.A,
            [transition_reply()],
            {"answer_sheet": None},
        )
        self.addCleanup(td.cleanup)
        self.assertEqual(self.value(out, result), (4, "preserve"))


if __name__ == "__main__":
    unittest.main()
