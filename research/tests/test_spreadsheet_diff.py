import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from adapters.spreadsheet_diff import diff_workbooks
from fulfilment import Scope


class SpreadsheetDiffOracleTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); root = Path(self.temp.name)
        self.before = root / "before.xlsx"; self.after = root / "after.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws["A1"] = ""; ws["A2"] = 0.10000000000000142
        ws["B1"] = "=A2*2"; ws["C1"] = "styled"
        ws["C1"].fill = PatternFill("solid", fgColor="FF0000")
        ws.merge_cells("D1:E1"); wb.create_sheet("Other")["A1"] = "preserve"
        wb.properties.title = "original"; wb.save(self.before); wb.close()
        shutil.copy2(self.before, self.after)

    def tearDown(self): self.temp.cleanup()

    def mutate(self, function):
        wb = load_workbook(self.after); function(wb); wb.save(self.after); wb.close()
        return diff_workbooks(self.before, self.after, (Scope("workbook/Data/F1"),))

    def test_float_roundtrip_is_conservatively_reported_until_dependency_equivalence_proven(self):
        def change(wb):
            wb["Data"]["A1"] = None
            wb["Data"]["A2"] = 0.10000000000001
        result = self.mutate(change)
        self.assertIn(Scope("workbook/Data/A2"), result.semantic_scopes)
        self.assertTrue(result.representation_parts)

    def test_formula_style_merge_metadata_and_other_sheet_changes_are_semantic(self):
        cases = {
            "formula": lambda wb: setattr(wb["Data"]["B1"], "value", "=A2*3"),
            "style": lambda wb: setattr(wb["Data"]["C1"], "fill", PatternFill("solid", fgColor="00FF00")),
            "merge": lambda wb: wb["Data"].unmerge_cells("D1:E1"),
            "metadata": lambda wb: setattr(wb.properties, "title", "changed"),
            "other-sheet": lambda wb: setattr(wb["Other"]["A1"], "value", "changed"),
        }
        for name, mutation in cases.items():
            with self.subTest(name=name):
                shutil.copy2(self.before, self.after)
                result = self.mutate(mutation)
                self.assertTrue(result.semantic_scopes)
                self.assertEqual(result.evaluator_visible_scopes, ())

    def test_evaluator_visible_projection_contains_only_declared_answer_scope(self):
        wb = load_workbook(self.after); wb["Data"]["F1"] = 9; wb["Other"]["A1"] = "changed"
        wb.save(self.after); wb.close()
        result = diff_workbooks(self.before, self.after, (Scope("workbook/Data/F1"),))
        self.assertEqual(result.evaluator_visible_scopes, (Scope("workbook/Data/F1"),))
        self.assertIn(Scope("workbook/Other/A1"), result.semantic_scopes)


if __name__ == "__main__": unittest.main()
