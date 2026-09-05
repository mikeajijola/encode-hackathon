import asyncio
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

RESEARCH = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RESEARCH))
sys.path.insert(0, str(RESEARCH / "baseline"))

from baseline.common import SpreadsheetAnswer, parse_answer, predict_task, prepare_out_dir, write_output
from evaluate import score, summarise


def workbook(path: Path, values: dict[str, object]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Input"
    for coordinate, value in values.items():
        ws[coordinate] = value
    wb.save(path)


class BaselineSemanticTests(unittest.TestCase):
    def test_parser_accepts_wrapped_json_and_strips_thinking(self):
        answer = parse_answer('<think>ignored</think> prose ```json\n{"cells":[{"cell":"B2","value":7}]}\n```')
        self.assertEqual(answer.cells[0].value, 7)

    def test_writer_mutates_only_declared_answer_cells(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            initial, output = root / "initial.xlsx", root / "output.xlsx"
            workbook(initial, {"A1": "preserve", "B2": 1, "C3": "preserve"})
            task = {"init_xlsx": str(initial), "answer_position": "B2", "answer_sheet": "Input"}
            answer = SpreadsheetAnswer.model_validate({"cells": [{"cell": "A1", "value": "bad"}, {"cell": "B2", "value": 9}]})
            write_output(task, answer, output)
            wb = openpyxl.load_workbook(output)
            self.assertEqual(wb["Input"]["B2"].value, 9)
            self.assertEqual(wb["Input"]["A1"].value, "preserve")
            self.assertEqual(wb["Input"]["C3"].value, "preserve")


class BaselineTraceTests(unittest.TestCase):
    def test_failed_call_falls_back_to_initial_and_emits_trace(self):
        async def fail(_prompt):
            raise RuntimeError("fixture failure")

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            initial = root / "initial.xlsx"
            workbook(initial, {"A1": "unchanged"})
            out = root / "run"
            prepare_out_dir(out)
            task = {
                "id": "fixture-1", "instruction": "change it", "init_xlsx": str(initial),
                "answer_position": "A1", "answer_sheet": "Input",
            }
            status = asyncio.run(predict_task(fail, "fixture-model", task, out))
            self.assertTrue(status.startswith("error:"))
            trace = json.loads((out / "traces" / "fixture-1.jsonl").read_text())
            prediction = json.loads((out / "predictions.jsonl").read_text())
            self.assertIn("RuntimeError", trace["error"])
            self.assertEqual(trace["step"], 1)
            self.assertEqual(prediction["status"], status)
            self.assertEqual(openpyxl.load_workbook(out / prediction["output"])["Input"]["A1"].value, "unchanged")


class EvaluatorStructuralTests(unittest.TestCase):
    def test_missing_task_counts_as_zero_cells_and_failure(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            golden = root / "golden.xlsx"
            workbook(golden, {"A1": 10, "A2": 20})
            task = {"id": "fixture", "instruction_type": "Cell-Level Manipulation", "golden_xlsx": str(golden),
                    "answer_position": "A1:A2", "answer_sheet": "Input"}
            summary, items = score([], [task], recalc=False)
            self.assertEqual(items[0]["status"], "missing")
            self.assertEqual(summary["pass_rate"], 0.0)
            self.assertEqual(summary["cell_accuracy"], 0.0)

    def test_error_rows_count_as_zero_in_cell_accuracy(self):
        items = [
            {"status": "graded", "cells": 2, "correct": 2, "pass": True, "type": "Cell-Level Manipulation"},
            {"status": "error", "cells": 2, "correct": 0, "pass": False, "type": "Cell-Level Manipulation"},
        ]
        self.assertEqual(summarise(items)["cell_accuracy"], 0.5)


if __name__ == "__main__":
    unittest.main()
