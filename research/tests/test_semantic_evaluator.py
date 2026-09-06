import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from adapters.spreadsheet import KIND
from experiment.runner import Arm, ExperimentRunner, ModelReply, RunConfig, Task
from services.spreadsheet import SpreadsheetServices
from fulfilment import EvidenceStore
from tests.contract_fixtures import capability_records, contract_reply, transition_reply


class RoleProvider:
    def __init__(self, replies): self.replies = list(replies); self.prompts = []
    def complete(self, prompt, *, model, temperature):
        self.prompts.append(prompt)
        return ModelReply(self.replies.pop(0), 3, 2, .01, f"role-{len(self.prompts)}")


def config(arm):
    return RunConfig("semantic", arm, "same-model", "pinned-v1", 0, 100, 8, 10000, 1,
                     "sha256:x", "none", "fixed", {"transport": 0})


class IsolatedSemanticEvaluatorTest(unittest.TestCase):
    def run_arm(self, arm, replies, service_type=SpreadsheetServices, context_extra=None):
        td = tempfile.TemporaryDirectory(); root = Path(td.name)
        artifact = root / "initial.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Data"; ws["A1"] = 2; ws["B1"] = None
        wb.save(artifact); wb.close()
        context = {"answer_sheet": "Data", "answer_position": "B1", "expected_type": "number",
                   "workbook_observation": {"facts": {"cells": [
                       {"selector": "Data!A1", "value": 2, "data_type": "n"},
                       {"selector": "Data!B1", "value": None, "data_type": "n"}]},
                       "truncation": {"truncated": False}}}
        context.update(context_extra or {})
        provider = RoleProvider(replies); out = root / "out"
        result = ExperimentRunner(config(arm), service_type(), provider, out).run(
            [Task("t", "put twice A1 in B1", artifact, KIND, context, capability_records())])[0]
        return td, out, provider, result

    def value(self, out, result):
        wb = load_workbook(out / result["output"]); value = wb["Data"]["B1"].value; wb.close(); return value

    def test_c_evaluates_wrong_once_and_never_repairs(self):
        replies = [
            contract_reply("B1 equals twice A1"),
            transition_reply(inputs={"writes": [{"selector": "Data!B1", "value": 3}]}),
            '{"verdict":"fail","expected_state":{"Data!B1":4},"rationale":"three is not twice two","confidence":0.99}',
        ]
        td, out, provider, result = self.run_arm(Arm.C, replies); self.addCleanup(td.cleanup)
        self.assertEqual(self.value(out, result), 3)
        self.assertFalse(result["terminal_eval"]["passed"])
        self.assertEqual(result["usage"]["model_calls"], 3)
        self.assertEqual(result["usage"]["tokens"], 15)
        c_evidence = EvidenceStore(out / "events" / "t.broker.jsonl").verify()
        self.assertEqual(c_evidence[-1].event_type, "termination_decision")
        self.assertFalse(c_evidence[-1].payload["fulfilled"])
        traces = [json.loads(line) for line in (out / "traces" / "t.jsonl").read_text().splitlines()]
        self.assertEqual([row["purpose"] for row in traces], ["contract", "action_generation", "independent_evaluation"])
        self.assertEqual(traces[-1]["provider_request_id"], "role-3")

    def test_d_acts_first_then_repairs_from_isolated_failed_verdict(self):
        replies = [
            contract_reply("B1 equals twice A1"),
            transition_reply(inputs={"writes": [{"selector": "Data!B1", "value": 3}]}),
            '{"verdict":"fail","expected_state":{"Data!B1":4},"rationale":"expected four","confidence":1}',
            transition_reply(),
            '{"verdict":"pass","expected_state":{"Data!B1":4},"rationale":"four is twice two","confidence":1}',
        ]
        td, out, provider, result = self.run_arm(Arm.D, replies); self.addCleanup(td.cleanup)
        self.assertEqual(self.value(out, result), 4)
        self.assertEqual(result["status"], "fulfilled")
        checkpoint = out / result["first_mutation_output"]
        wb = load_workbook(checkpoint); self.assertEqual(wb["Data"]["B1"].value, 3); wb.close()
        checkpoint_rows = [json.loads(line) for line in
                           (out / "first_mutation_predictions.jsonl").read_text().splitlines()]
        self.assertEqual(checkpoint_rows[0]["output"], result["first_mutation_output"])
        self.assertEqual(result["usage"]["model_calls"], 5)
        self.assertEqual(result["usage"]["tokens"], 25)
        traces = [json.loads(line) for line in (out / "traces" / "t.jsonl").read_text().splitlines()]
        self.assertEqual([row["purpose"] for row in traces], ["contract", "action_generation",
                         "independent_evaluation", "action_generation", "independent_evaluation"])
        # The first action occurs before the first independent-evaluation call.
        events = [json.loads(line) for line in (out / "events" / "t.jsonl").read_text().splitlines()]
        self.assertTrue(any(e["event_type"] == "semantic_evaluator_verdict" for e in events))
        evidence = EvidenceStore(out / "events" / "t.broker.jsonl").verify()
        self.assertEqual(evidence[-1].event_type, "termination_decision")
        self.assertTrue(evidence[-1].payload["fulfilled"])
        self.assertEqual(evidence[-1].payload["artifact"]["hash"], result["output_artifact_hash"])
        self.assertTrue({"accepted_contract", "observation", "eval_result", "termination_decision"}
                        <= {event.event_type for event in evidence})
        evaluator_prompts = [p for p in provider.prompts if "independent_semantic_evaluator" in p]
        self.assertEqual(len(evaluator_prompts), 2)
        self.assertTrue(all("writes" not in p and "golden" not in p.lower() for p in evaluator_prompts))
        repair_prompt = provider.prompts[3]
        self.assertIn("state_discrepancy", repair_prompt)
        self.assertNotIn("independent_semantic_evaluator", repair_prompt)

    def test_semantic_pass_with_truncated_source_evidence_is_unknown(self):
        replies = [
            contract_reply("B1 equals twice A1"),
            transition_reply(),
            '{"verdict":"pass","expected_state":{"Data!B1":4},"rationale":"looks correct","confidence":1}',
        ]
        truncated = {"workbook_observation": {"facts": {"cells": [
            {"selector": "Data!A1", "value": 2, "data_type": "n"}
        ]}, "truncation": {"truncated": True, "omitted_nonempty_cells": 10}}}
        td, out, provider, result = self.run_arm(Arm.D, replies, context_extra=truncated)
        self.addCleanup(td.cleanup)
        self.assertEqual(self.value(out, result), 4)
        self.assertEqual(result["internal_status"], "UNKNOWN")
        semantic = next(item for item in result["terminal_eval"]["details"]["evals"]
                        if item["eval_id"] == "semantic")
        self.assertEqual(semantic["status"], "uncertain")
        self.assertEqual(semantic["details"]["epistemic_reason"], "source_observation_truncated")

    def test_malformed_evaluator_response_is_uncertain_and_never_passes(self):
        replies = [contract_reply("B1 equals twice A1"),
                   transition_reply(),
                   '{"verdict":"pass"}']
        td, out, provider, result = self.run_arm(Arm.D, replies); self.addCleanup(td.cleanup)
        self.assertEqual(result["status"], "unknown:evaluation_uncertain")
        self.assertEqual(result["internal_status"], "UNKNOWN")
        self.assertFalse(result["terminal_eval"]["passed"])
        semantic = next(e for e in result["terminal_eval"]["details"]["evals"] if e["eval_id"] == "semantic")
        self.assertEqual(semantic["status"], "uncertain")
        self.assertIn("malformed", semantic["message"])

    def test_explicit_uncertain_verdict_prevents_fulfilled(self):
        replies = [contract_reply("B1 equals twice A1"),
                   transition_reply(),
                   '{"verdict":"uncertain","expected_state":null,"rationale":"insufficient source facts","confidence":0.2}']
        td, out, provider, result = self.run_arm(Arm.D, replies); self.addCleanup(td.cleanup)
        self.assertEqual(result["status"], "unknown:evaluation_uncertain")
        self.assertEqual(result["internal_status"], "UNKNOWN")
        self.assertFalse(EvidenceStore(out / "events" / "t.broker.jsonl").verify()[-1].payload["fulfilled"])

    def test_evaluator_provider_error_prevents_fulfilled(self):
        replies = [contract_reply("B1 equals twice A1"),
                   transition_reply()]
        td, out, provider, result = self.run_arm(Arm.D, replies); self.addCleanup(td.cleanup)
        self.assertEqual(result["status"], "unknown:evaluation_uncertain")
        self.assertEqual(result["internal_status"], "UNKNOWN")
        semantic = next(e for e in result["terminal_eval"]["details"]["evals"] if e["eval_id"] == "semantic")
        self.assertEqual(semantic["status"], "uncertain")
        self.assertIn("IndexError", semantic["message"])
        self.assertFalse(EvidenceStore(out / "events" / "t.broker.jsonl").verify()[-1].payload["fulfilled"])

    def test_generic_layer_has_no_spreadsheet_imports(self):
        fulfilment = Path(__file__).resolve().parents[1] / "fulfilment"
        source = "\n".join(path.read_text() for path in fulfilment.glob("*.py"))
        self.assertNotIn("openpyxl", source)
        self.assertNotIn("adapters.spreadsheet", source)

    def test_generic_repeated_transition_stops_no_progress(self):
        replies = [contract_reply("B1 equals twice A1"),
                   transition_reply(inputs={"writes": [{"selector": "Data!B1", "value": 3}]}),
                   '{"verdict":"fail","expected_state":{"Data!B1":4},"rationale":"expected four","confidence":1}',
                   transition_reply(inputs={"writes": [{"selector": "Data!B1", "value": 3}]})]
        td, out, provider, result = self.run_arm(Arm.D, replies); self.addCleanup(td.cleanup)
        self.assertEqual(result["status"], "unfulfilled:no_progress")
        evidence = EvidenceStore(out / "events" / "t.broker.jsonl").verify()
        self.assertEqual(evidence[-1].payload["reason"], "no_progress")
        self.assertEqual(sum(event.event_type == "capability_request" for event in evidence), 1)

    def test_generic_knowledge_gap_observes_then_acts(self):
        replies = [contract_reply("B1 equals twice A1"),
                   transition_reply(),
                   '{"verdict":"pass","expected_state":{"Data!B1":4},"rationale":"correct","confidence":1}']
        td, out, provider, result = self.run_arm(
            Arm.D, replies, context_extra={"test_omit_target_once": True})
        self.addCleanup(td.cleanup)
        self.assertEqual(result["status"], "fulfilled")
        evidence = EvidenceStore(out / "events" / "t.broker.jsonl").verify()
        transitions = [event for event in evidence if event.event_type == "transition_decision"]
        self.assertEqual(transitions[0].payload["kind"], "observe")
        self.assertEqual(transitions[1].payload["kind"], "capability")

    def test_generic_broker_rejects_d_scope_violation(self):
        replies = [contract_reply("B1 equals twice A1"),
                   transition_reply(inputs={"writes": [{"selector": "Data!A1", "value": 99}]})]
        td, out, provider, result = self.run_arm(Arm.D, replies); self.addCleanup(td.cleanup)
        self.assertEqual(result["status"], "unfulfilled:no_safe_transition")
        self.assertEqual(self.value(out, result), None)
        evidence = EvidenceStore(out / "events" / "t.broker.jsonl").verify()
        self.assertFalse(evidence[-1].payload["fulfilled"])

    def test_a_and_b_never_invoke_evaluation_role(self):
        cases = {
            Arm.A: [transition_reply()],
            Arm.B: [contract_reply("B1 equals twice A1"), transition_reply()],
        }
        for arm, replies in cases.items():
            with self.subTest(arm=arm):
                td, out, provider, result = self.run_arm(arm, replies); self.addCleanup(td.cleanup)
                self.assertFalse(any("independent_semantic_evaluator" in p for p in provider.prompts))
                self.assertEqual(result["usage"]["model_calls"], 1 if arm is Arm.A else 2)


if __name__ == "__main__": unittest.main()
