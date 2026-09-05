import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from backends.spreadsheetbench import (
    OpenRouterProvider, ProviderConfigurationError, ScriptedProvider,
    bounded_workbook_context, factory, load_runtime_tasks,
)
from experiment.runner import Arm, ExperimentRunner, RunConfig
from experiment.cli import main as cli_main


class SpreadsheetBenchBackendTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        folder = self.root / "spreadsheet" / "x"; folder.mkdir(parents=True)
        self.init = folder / "1_x_init.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Data"
        ws["A1"] = 2; ws["A2"] = "=A1*2"; ws["A2"].number_format = "0.00"
        for row in range(3, 15): ws.cell(row, 1).value = row
        wb.save(self.init); wb.close()
        # A golden-looking sentinel exists but runtime loading must never open it.
        self.golden = folder / "1_x_golden.xlsx"; self.golden.write_bytes(b"must-not-open")
        dataset = [{"id": "x", "instruction": "put twice A1 in B1", "spreadsheet_path": "spreadsheet/x",
                    "instruction_type": "Cell-Level Manipulation", "answer_sheet": "Data",
                    "answer_position": "B1", "data_position": "A1:B14",
                    "golden_xlsx": str(self.golden), "expected_answer": 4}]
        (self.root / "dataset.json").write_text(json.dumps(dataset))
        self.selection = self.root / "selection.json"; self.selection.write_text(json.dumps({"task_ids": ["x"]}))

    def tearDown(self): self.temp.cleanup()

    def raw(self, arm="A", replies=None):
        return {"backend": "backends.spreadsheetbench:factory", "backend_config": {
            "dataset_dir": str(self.root), "selection_manifest": str(self.selection),
            "provider": {"type": "scripted", "replies": replies or ['{"writes":[{"selector":"Data!B1","value":4}]}']},
            "context_limits": {"max_cells": 5, "max_chars": 2000}},
            "run_config": {"experiment_id": "e", "arm": arm, "model": "pinned", "model_version": "v1",
            "temperature": 0, "max_tokens": 100, "max_actions": 5, "max_wall_time_ms": 10000,
            "max_cost": 1, "environment_image_digest": "sha256:x", "recalculation_engine": "none",
            "scorer_commit": "fixed", "retry_policy": {"transport": 0}}}

    def test_loader_exposes_only_init_metadata_and_bounded_canonical_context(self):
        real_load = __import__("backends.spreadsheetbench", fromlist=["load_workbook"]).load_workbook
        opened = []
        def recording(path, *args, **kwargs):
            opened.append(str(path)); return real_load(path, *args, **kwargs)
        with patch("backends.spreadsheetbench.load_workbook", side_effect=recording):
            tasks = load_runtime_tasks(self.root, selected_ids={"x"}, max_cells=5, max_chars=2000)
        self.assertEqual(opened, [str(self.init)])
        task = tasks[0]
        serialized = json.dumps({"context": task.context, "artifact": str(task.artifact)}, default=str)
        self.assertNotIn("golden", serialized.lower())
        self.assertNotIn(str(self.golden), serialized)
        observation = task.context["workbook_observation"]
        self.assertTrue(observation["truncation"]["truncated"])
        self.assertEqual(observation["truncation"]["serialized_cells"], 5)
        formula = next(cell for cell in observation["facts"]["cells"] if cell["selector"] == "Data!A2")
        self.assertEqual(formula["value"], "=A1*2")
        self.assertEqual(formula["data_type"], "f")
        self.assertEqual(formula["style"]["number_format"], "0.00")

    def test_factory_selection_capabilities_and_context_are_arm_invariant(self):
        signatures = []
        for arm in "ABCD":
            services, provider, tasks = factory(self.raw(arm))
            signatures.append((tasks[0].context, tasks[0].capability_manifests))
            self.assertIsInstance(provider, ScriptedProvider)
        self.assertTrue(all(item == signatures[0] for item in signatures[1:]))
        self.assertEqual(len(signatures[0][1]), 6)

    def test_arm_a_remains_one_model_call_and_receives_bounded_observation(self):
        raw = self.raw("A")
        services, provider, tasks = factory(raw)
        fields = {key: value for key, value in raw["run_config"].items() if key != "arm"}
        out = self.root / "out"
        results = ExperimentRunner(RunConfig(arm=Arm.A, **fields), services, provider, out).run(tasks)
        self.assertEqual(results[0]["usage"]["model_calls"], 1)
        self.assertEqual(results[0]["usage"]["actions"], 1)
        self.assertEqual(len(provider.prompts), 1)
        self.assertIn("workbook_observation", provider.prompts[0])
        self.assertNotIn("golden", provider.prompts[0].lower())

    def test_contract_and_action_prompts_receive_same_bounded_context(self):
        replies = ['{"description":"B1 is twice A1"}', '{"writes":[{"selector":"Data!B1","value":4}]}']
        raw = self.raw("B", replies)
        services, provider, tasks = factory(raw)
        fields = {key: value for key, value in raw["run_config"].items() if key != "arm"}
        ExperimentRunner(RunConfig(arm=Arm.B, **fields), services, provider, self.root / "out-b").run(tasks)
        self.assertEqual(len(provider.prompts), 2)
        for prompt in provider.prompts:
            self.assertIn("workbook_observation", prompt)
            self.assertIn("=A1*2", prompt)
            self.assertNotIn("golden", prompt.lower())

    def test_openrouter_missing_key_and_nonzero_temperature_are_typed(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaisesRegex(ProviderConfigurationError, "OPENROUTER_API_KEY"):
                OpenRouterProvider().complete("x", model="pinned", temperature=0)
        raw = self.raw(); raw["run_config"]["temperature"] = .2
        with self.assertRaisesRegex(ProviderConfigurationError, "temperature 0"):
            factory(raw)

    def test_unknown_selection_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "unknown task"):
            load_runtime_tasks(self.root, selected_ids={"absent"}, max_cells=5, max_chars=1000)

    def test_cli_loads_factory_manifest_and_selected_task(self):
        manifest = self.root / "run.json"; manifest.write_text(json.dumps(self.raw("A")))
        out = self.root / "cli-out"
        with patch.object(sys, "argv", ["experiment", "--manifest", str(manifest), "--out-dir", str(out)]):
            cli_main()
        prediction = json.loads((out / "predictions.jsonl").read_text())
        self.assertEqual(prediction["id"], "x")
        run_manifest = json.loads((out / "run_manifest.json").read_text())
        self.assertEqual(run_manifest["model"], "pinned")
        self.assertEqual(run_manifest["temperature"], 0)


if __name__ == "__main__": unittest.main()
