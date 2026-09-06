"""Targeted OOXML cell transplantation for low-drift reconstruction."""

from __future__ import annotations

from copy import copy, deepcopy
from pathlib import Path
import xml.etree.ElementTree as ET
import re
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter

from fulfilment.models import Scope


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", MAIN)
ET.register_namespace("r", REL)


def _sheet_parts(archive: ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.attrib["Id"]: item.attrib["Target"]
               for item in relations.findall(f"{{{PKG_REL}}}Relationship")}
    result = {}
    for sheet in workbook.findall(f"{{{MAIN}}}sheets/{{{MAIN}}}sheet"):
        target = targets[sheet.attrib[f"{{{REL}}}id"]].lstrip("/")
        result[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
    return result


def patch_cells_from_candidate(original: Path, candidate: Path, output: Path,
                               scopes: tuple[Scope, ...]) -> None:
    """Copy only authorized cell XML from candidate into the original package."""
    selected: dict[str, set[str]] = {}
    for scope in scopes:
        pieces = scope.resource.split("/", 2)
        if len(pieces) != 3 or pieces[0] != "workbook" or pieces[2].startswith("__"):
            raise ValueError(f"OOXML patch requires exact cell scope: {scope.resource}")
        selected.setdefault(pieces[1], set()).add(pieces[2])
    with ZipFile(original) as source, ZipFile(candidate) as changed:
        source_parts, changed_parts = _sheet_parts(source), _sheet_parts(changed)
        replacements: dict[str, bytes] = {}
        for sheet, coordinates in selected.items():
            if sheet not in source_parts or sheet not in changed_parts:
                raise ValueError(f"worksheet not found: {sheet}")
            source_xml = ET.fromstring(source.read(source_parts[sheet]))
            changed_xml = ET.fromstring(changed.read(changed_parts[sheet]))
            source_data = source_xml.find(f"{{{MAIN}}}sheetData")
            changed_data = changed_xml.find(f"{{{MAIN}}}sheetData")
            assert source_data is not None and changed_data is not None
            changed_cells = {cell.attrib["r"]: cell for cell in changed_data.iter(f"{{{MAIN}}}c")}
            for coordinate in coordinates:
                replacement = changed_cells.get(coordinate)
                existing = next((cell for cell in source_data.iter(f"{{{MAIN}}}c")
                                 if cell.attrib.get("r") == coordinate), None)
                if existing is not None:
                    row = next(row for row in source_data.findall(f"{{{MAIN}}}row") if existing in list(row))
                    index = list(row).index(existing)
                    if replacement is None:
                        row.remove(existing)
                    else:
                        row.remove(existing); row.insert(index, deepcopy(replacement))
                elif replacement is not None:
                    row_number = int("".join(ch for ch in coordinate if ch.isdigit()))
                    row = next((item for item in source_data.findall(f"{{{MAIN}}}row")
                                if int(item.attrib["r"]) == row_number), None)
                    if row is None:
                        row = ET.Element(f"{{{MAIN}}}row", {"r": str(row_number)})
                        source_data.append(row)
                    row.append(deepcopy(replacement))
            replacements[source_parts[sheet]] = ET.tostring(source_xml, encoding="utf-8",
                                                              xml_declaration=True)
        with ZipFile(output, "w") as target:
            for info in source.infolist():
                target.writestr(info, replacements.get(info.filename, source.read(info.filename)))


def patch_cell_values_raw(original: Path, candidate: Path, output: Path,
                          scopes: tuple[Scope, ...]) -> None:
    """Value-only reconstruction preserving every unselected XML byte.

    Current mutators write values/formulas only. Styles, relationships, layout,
    metadata and all other package members come from the original. Candidate
    style changes are not part of value-only capabilities and are not copied.
    Unsupported namespace layouts fail closed; no ordinary-save fallback.
    """
    selected: dict[str, set[str]] = {}
    for scope in scopes:
        parts = scope.resource.split("/", 2)
        if len(parts) != 3 or parts[0] != "workbook" or not re.fullmatch(r"[A-Z]+[1-9][0-9]*", parts[2]):
            raise ValueError("raw reconstruction requires exact cell scopes")
        selected.setdefault(parts[1], set()).add(parts[2])
    before_book = load_workbook(original)
    candidate_book = load_workbook(candidate)
    try:
        with ZipFile(original) as source, ZipFile(candidate) as changed:
            source_parts, changed_parts = _sheet_parts(source), _sheet_parts(changed)
            replacements = {}
            for sheet, coordinates in selected.items():
                raw = source.read(source_parts[sheet])
                if not re.search(rb"<worksheet(?:\s|>)", raw):
                    raise ValueError("unsupported prefixed worksheet XML")
                changed_tree = ET.fromstring(changed.read(changed_parts[sheet]))
                cells = {cell.attrib["r"]: cell for cell in changed_tree.iter(f"{{{MAIN}}}c")}
                for tree in (ET.fromstring(raw), changed_tree):
                    for cell in tree.iter(f"{{{MAIN}}}c"):
                        formula = cell.find(f"{{{MAIN}}}f")
                        if formula is not None and formula.attrib.get("ref"):
                            lo_col, lo_row, hi_col, hi_row = range_boundaries(formula.attrib["ref"])
                            group = {f"{get_column_letter(col)}{row}" for row in range(lo_row, hi_row + 1)
                                     for col in range(lo_col, hi_col + 1)}
                            if coordinates & group and not group <= coordinates:
                                raise ValueError("formula group crosses authorized scope")
                for coordinate in sorted(coordinates, key=coordinate_to_tuple):
                    old_cell, new_cell = before_book[sheet][coordinate], candidate_book[sheet][coordinate]
                    pattern = rb'<c\b(?=[^>]*\br="' + coordinate.encode() + rb'")[^>]*?(?:/>|>.*?</c>)'
                    match = re.search(pattern, raw, flags=re.S)
                    node = deepcopy(cells.get(coordinate))
                    replacement = b""
                    if node is not None:
                        # Candidate style/shared-string IDs belong to its package.
                        node.attrib.pop("s", None)
                        if match:
                            style = re.search(rb'\bs="([0-9]+)"', match.group())
                            if style:
                                node.attrib["s"] = style.group(1).decode()
                        if node.attrib.get("t") == "s":
                            node.attrib["t"] = "inlineStr"
                            for child in list(node):
                                node.remove(child)
                            inline = ET.SubElement(node, f"{{{MAIN}}}is")
                            text = ET.SubElement(inline, f"{{{MAIN}}}t")
                            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                            text.text = new_cell.value
                        formula = node.find(f"{{{MAIN}}}f")
                        if formula is not None and formula.attrib.get("t") == "shared":
                            raise ValueError("shared formula requires group-aware reconstruction")
                        replacement = ET.tostring(node, encoding="utf-8")
                    if match:
                        raw = raw[:match.start()] + replacement + raw[match.end():]
                    elif replacement:
                        row_number = coordinate_to_tuple(coordinate)[0]
                        row_pattern = rb'<row\b(?=[^>]*\br="' + str(row_number).encode() + rb'")[^>]*?(?:/>|>.*?</row>)'
                        row_match = re.search(row_pattern, raw, flags=re.S)
                        if row_match:
                            row = row_match.group()
                            if row.endswith(b"/>"):
                                row = row[:-2] + b">" + replacement + b"</row>"
                            else:
                                following = next((m for m in re.finditer(rb'<c\b[^>]*\br="([A-Z]+[0-9]+)"', row)
                                                  if coordinate_to_tuple(m.group(1).decode())[1] > coordinate_to_tuple(coordinate)[1]), None)
                                index = following.start() if following else row.rfind(b"</row>")
                                row = row[:index] + replacement + row[index:]
                            raw = raw[:row_match.start()] + row + raw[row_match.end():]
                        else:
                            row = b'<row r="' + str(row_number).encode() + b'">' + replacement + b'</row>'
                            following = next((m for m in re.finditer(rb'<row\b[^>]*\br="([0-9]+)"', raw)
                                              if int(m.group(1)) > row_number), None)
                            index = following.start() if following else raw.find(b"</sheetData>")
                            if index < 0:
                                raise ValueError("unsupported empty sheetData encoding")
                            raw = raw[:index] + row + raw[index:]
                ET.fromstring(raw)
                replacements[source_parts[sheet]] = raw
            with ZipFile(output, "w") as target:
                for info in source.infolist():
                    target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
    finally:
        before_book.close()
        candidate_book.close()
