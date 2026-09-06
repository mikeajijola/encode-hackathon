"""Independent semantic, package, and evaluator-visible XLSX diffs."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from hashlib import sha256
import math
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.formula import ArrayFormula

from fulfilment.models import Scope


@dataclass(frozen=True)
class SpreadsheetDiff:
    semantic_scopes: tuple[Scope, ...]
    representation_parts: tuple[str, ...]
    evaluator_visible_scopes: tuple[Scope, ...]
    semantic_changes: tuple[dict[str, Any], ...]
    attempted_semantic_scopes: tuple[Scope, ...] = ()
    reconstruction_applied: bool = False

    def evidence(self) -> dict[str, Any]:
        return {
            "semantic_scopes": [item.resource for item in self.semantic_scopes],
            "representation_parts": list(self.representation_parts),
            "evaluator_visible_scopes": [item.resource for item in self.evaluator_visible_scopes],
            "semantic_changes": list(self.semantic_changes),
            "attempted_semantic_scopes": [item.resource for item in self.attempted_semantic_scopes],
            "reconstruction_applied": self.reconstruction_applied,
        }


def _value(value: Any) -> Any:
    if isinstance(value, ArrayFormula):
        return ("array_formula", value.text, value.ref)
    if isinstance(value, datetime):
        return ("excel_datetime", float(to_excel(value)))
    if isinstance(value, date):
        return ("excel_date", float(to_excel(value)))
    if isinstance(value, time):
        return ("excel_time", value.isoformat())
    return value


def _equivalent(left: Any, right: Any) -> bool:
    left, right = _value(left), _value(right)
    if left is None and right is None:
        return True
    if isinstance(left, (int, float)) and not isinstance(left, bool) and \
       isinstance(right, (int, float)) and not isinstance(right, bool):
        return left == right
    return type(left) is type(right) and left == right


def _cell_state(cell) -> dict[str, Any]:
    value = _value(cell.value)
    # A missing value has no scalar type. Empty string remains distinguishable:
    # dependent formulas can distinguish it from an absent cell.
    data_type = None if value is None else cell.data_type
    return {
        "value": value,
        "data_type": data_type,
        "number_format": cell.number_format,
        "style": tuple(str(getattr(cell, key)) for key in
                       ("font", "fill", "border", "alignment", "protection")),
        "hyperlink": None if cell.hyperlink is None else cell.hyperlink.target,
        "comment": None if cell.comment is None else (cell.comment.text, cell.comment.author),
    }


def semantic_snapshot(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False, read_only=False)
    result: dict[str, Any] = {}
    result["workbook/__structure__/sheets"] = tuple(
        (ws.title, ws.sheet_state) for ws in wb.worksheets)
    result["workbook/__metadata__/properties"] = (
        wb.properties.title, wb.properties.subject, wb.properties.creator,
        wb.properties.description, wb.properties.keywords, wb.properties.category,
    )
    result["workbook/__structure__/defined_names"] = tuple(
        sorted((name, str(item)) for name, item in wb.defined_names.items()))
    for ws in wb.worksheets:
        result[f"workbook/{ws.title}/__merges__"] = tuple(sorted(str(x) for x in ws.merged_cells.ranges))
        result[f"workbook/{ws.title}/__dimensions__"] = (
            tuple(sorted((key, value.width, value.hidden) for key, value in ws.column_dimensions.items())),
            tuple(sorted((key, value.height, value.hidden) for key, value in ws.row_dimensions.items())),
        )
        for row in ws.iter_rows():
            for cell in row:
                state = _cell_state(cell)
                if state["value"] is not None or cell.has_style or cell.comment or cell.hyperlink:
                    result[f"workbook/{ws.title}/{cell.coordinate}"] = state
    wb.close()
    return result


def _state_equivalent(left: Any, right: Any) -> bool:
    if not isinstance(left, dict) or not isinstance(right, dict):
        return left == right
    if not _equivalent(left.get("value"), right.get("value")):
        return False
    return all(left.get(key) == right.get(key)
               for key in ("data_type", "number_format", "style", "hyperlink", "comment"))


def package_snapshot(path: Path) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {name: sha256(archive.read(name)).hexdigest() for name in archive.namelist()}


def diff_workbooks(before: Path, after: Path,
                   evaluator_scopes: tuple[Scope, ...] = ()) -> SpreadsheetDiff:
    left, right = semantic_snapshot(before), semantic_snapshot(after)
    changes = diff_semantic_snapshots(left, right)
    semantic = tuple(Scope(item["scope"]) for item in changes)
    # Cached-value projection only. Actual official visibility is established
    # offline after LibreOffice recalculation, which can propagate dependencies.
    evaluator = []
    before_cached = load_workbook(before, data_only=True)
    after_cached = load_workbook(after, data_only=True)
    try:
        for scope in evaluator_scopes:
            _, sheet, coord = scope.resource.split("/", 2)
            if sheet not in before_cached or sheet not in after_cached:
                evaluator.append(scope)
            elif not evaluator_values_equal(before_cached[sheet][coord].value,
                                            after_cached[sheet][coord].value):
                evaluator.append(scope)
    finally:
        before_cached.close(); after_cached.close()
    left_package, right_package = package_snapshot(before), package_snapshot(after)
    parts = tuple(sorted(name for name in left_package.keys() | right_package.keys()
                         if left_package.get(name) != right_package.get(name)))
    return SpreadsheetDiff(semantic, parts, tuple(evaluator), changes)


def evaluator_values_equal(left: Any, right: Any) -> bool:
    """Gold-free cached-value comparator; offline scorer remains unchanged."""
    def normalize(value):
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return round(float(value), 2)
        if isinstance(value, datetime):
            return round(float(to_excel(value)), 0)
        if isinstance(value, time):
            return str(value)[:-3]
        if isinstance(value, str):
            try:
                return round(float(value), 2)
            except ValueError:
                return value
        return value
    left, right = normalize(left), normalize(right)
    return (left in (None, "") and right in (None, "")) or (type(left) is type(right) and left == right)


def diff_semantic_snapshots(left: dict[str, Any], right: dict[str, Any]) -> tuple[dict[str, Any], ...]:
    changes = []
    for key in sorted(left.keys() | right.keys()):
        if _state_equivalent(left.get(key), right.get(key)):
            continue
        changes.append({"scope": key, "before": repr(left.get(key)), "after": repr(right.get(key))})
    return tuple(changes)
