"""Minimal XLSX adapter. No spreadsheet types escape through broker interfaces."""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import tempfile
from dataclasses import replace
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils.cell import get_column_letter, range_boundaries

from fulfilment.models import CapabilityEffect, CapabilityManifest, CapabilityRequest, CapabilityResult, Scope
from adapters.spreadsheet_diff import diff_workbooks
from adapters.spreadsheet_ooxml import patch_cell_values_raw

KIND = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VERSION = "1.0.0"


class SelectorError(ValueError):
    pass


def parse_selector(selector: str) -> tuple[str, str]:
    if "!" not in selector:
        raise SelectorError("selector must be explicit Sheet!A1 or Sheet!A1:B2")
    sheet, coordinates = selector.rsplit("!", 1)
    sheet = sheet.strip("'")
    if not sheet or not coordinates:
        raise SelectorError("selector has an empty sheet or coordinate")
    try:
        range_boundaries(coordinates)
    except (TypeError, ValueError) as error:
        raise SelectorError(f"invalid cell range: {coordinates}") from error
    return sheet, coordinates.upper()


def scope_for(selector: str) -> Scope:
    sheet, coordinates = parse_selector(selector)
    return Scope(f"workbook/{sheet}/{coordinates}")


def expand_selector_scopes(selector: str) -> tuple[Scope, ...]:
    """Translate one compact selector into exact cell-level mutation scopes."""
    sheet, coordinates = parse_selector(selector)
    min_col, min_row, max_col, max_row = range_boundaries(coordinates)
    return tuple(Scope(f"workbook/{sheet}/{get_column_letter(column)}{row}")
                 for row in range(min_row, max_row + 1)
                 for column in range(min_col, max_col + 1))


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cells(ws, coordinates: str):
    min_col, min_row, max_col, max_row = range_boundaries(coordinates)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        yield from row


def _value(value: Any) -> Any:
    if isinstance(value, ArrayFormula):
        return value.text
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


class WorkbookSnapshots:
    """Broker snapshot hooks plus adapter-level changed-cell translation."""

    def __init__(self):
        self.before: dict[str, dict[str, Any]] = {}
        self.diffs: dict[str, list[dict[str, Any]]] = {}
        self.originals: dict[str, Path] = {}
        self.staged: dict[str, Path] = {}

    def before_mutation(self, request: CapabilityRequest) -> str:
        path = Path(request.artifact_id)
        self.before[request.id] = _snapshot_cells(path)
        return _hash(path)

    def after_mutation(self, request: CapabilityRequest) -> str:
        path = Path(request.artifact_id)
        after = _snapshot_cells(path)
        before = self.before.pop(request.id)
        keys = sorted(before.keys() | after.keys())
        self.diffs[request.id] = [
            {"scope": key, "before": before.get(key), "after": after.get(key)}
            for key in keys if before.get(key) != after.get(key)
        ]
        return _hash(path)

    def stage_mutation(self, request: CapabilityRequest) -> CapabilityRequest:
        original = Path(request.artifact_id)
        staged = original.with_name(f".{original.stem}.transaction-{request.id}{original.suffix}")
        shutil.copy2(original, staged)
        self.originals[request.id] = original
        self.staged[request.id] = staged
        return replace(request, artifact_id=str(staged))

    def inspect_mutation(self, request: CapabilityRequest, staged: CapabilityRequest):
        original, candidate = Path(request.artifact_id), Path(staged.artifact_id)
        attempted = diff_workbooks(original, candidate, request.requested_mutation_scope)
        # A mutator may have produced useful in-scope state while its serializer
        # also changed unrelated state. Rebuild from the exact original package,
        # transplanting only authorized cell nodes, then re-inspect what will be
        # committed. This adapter hook implements the artifact-specific repair;
        # authorization and transaction control remain generic in the broker.
        reconstructed = candidate.with_name(f".{candidate.stem}.reconstructed{candidate.suffix}")
        patch_cell_values_raw(original, candidate, reconstructed,
                              request.requested_mutation_scope)
        reconstructed.replace(candidate)
        committed = diff_workbooks(original, candidate, request.requested_mutation_scope)
        return replace(committed,
                       attempted_semantic_scopes=attempted.semantic_scopes,
                       reconstruction_applied=True)

    def commit_mutation(self, request: CapabilityRequest, staged: CapabilityRequest) -> None:
        Path(staged.artifact_id).replace(request.artifact_id)
        self.originals.pop(request.id, None); self.staged.pop(request.id, None)

    def rollback_mutation(self, request: CapabilityRequest, staged: CapabilityRequest) -> None:
        Path(staged.artifact_id).unlink(missing_ok=True)
        candidate = Path(staged.artifact_id)
        candidate.with_name(f".{candidate.stem}.reconstructed{candidate.suffix}").unlink(missing_ok=True)
        self.originals.pop(request.id, None); self.staged.pop(request.id, None)


def _snapshot_cells(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False, read_only=False)
    result = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    result[f"workbook/{ws.title}/{cell.coordinate}"] = _value(cell.value)
    wb.close()
    return result


class SpreadsheetCapability:
    def __init__(self, name: str):
        self.name = name

    def invoke(self, request: CapabilityRequest) -> CapabilityResult:
        path = Path(request.artifact_id)
        try:
            if self.name == "inspect_workbook":
                return self._inspect(request, path)
            if self.name == "write_cells":
                return self._write(request, path)
            if self.name == "copy_or_fill_formula":
                return self._copy_formula(request, path)
            if self.name == "recalculate":
                return self._recalculate(request, path)
            if self.name == "validate_workbook":
                return self._validate(request, path)
            if self.name == "render_range":
                return self._render(request, path)
            raise ValueError(f"unknown adapter capability {self.name}")
        except (OSError, ValueError, KeyError, SelectorError) as error:
            return self._result(request, False, {}, error=f"adapter_error:{type(error).__name__}:{error}")

    def _provenance(self, path: Path, selectors=()) -> dict:
        return {"adapter_version": VERSION, "artifact_sha256": _hash(path) if path.exists() else None,
                "selectors": list(selectors)}

    def _result(self, request, succeeded, output, scopes=(), selectors=(), error=None):
        return CapabilityResult(request.id, succeeded, output, tuple(scopes),
                                self._provenance(Path(request.artifact_id), selectors), error)

    def _inspect(self, request, path):
        if not isinstance(request.inputs["selectors"], (list, tuple)):
            raise ValueError("selectors must be an array")
        selectors = list(request.inputs["selectors"])
        wb = load_workbook(path, data_only=False, read_only=False)
        facts = []
        inspected = []
        for selector in selectors:
            sheet, coordinates = parse_selector(selector)
            if sheet not in wb.sheetnames:
                raise SelectorError(f"sheet not found: {sheet}")
            inspected.append(scope_for(selector).resource)
            for cell in _cells(wb[sheet], coordinates):
                facts.append({"selector": f"{sheet}!{cell.coordinate}", "value": _value(cell.value),
                              "data_type": cell.data_type, "number_format": cell.number_format})
        omitted = [f"workbook/{name}/*" for name in wb.sheetnames]
        output = {"observation": {"artifact_kind": KIND, "artifact_hash": _hash(path),
                  "facts": {"cells": facts, "sheet_names": wb.sheetnames},
                  "interpretations": {}, "inspected_scope": inspected, "omitted_scope": omitted}}
        wb.close()
        return self._result(request, True, output, selectors=selectors)

    def _write(self, request, path):
        if not isinstance(request.inputs["writes"], (list, tuple)) or not all(isinstance(item, dict) and "selector" in item for item in request.inputs["writes"]):
            raise ValueError("writes must be an array of objects with selectors")
        writes = list(request.inputs["writes"])
        selectors = [item["selector"] for item in writes]
        scopes = tuple(scope_for(s) for s in selectors)
        _require_requested(scopes, request.requested_mutation_scope)
        wb = load_workbook(path)
        for item in writes:
            sheet, coordinates = parse_selector(item["selector"])
            min_col, min_row, max_col, max_row = range_boundaries(coordinates)
            if (min_col, min_row) != (max_col, max_row):
                raise SelectorError("write selector must identify one cell")
            wb[sheet].cell(min_row, min_col).value = item.get("value")
        wb.save(path); wb.close()
        return self._result(request, True, {"written": selectors}, scopes, selectors)

    def _copy_formula(self, request, path):
        source = request.inputs["source"]
        target = request.inputs["target"]
        scopes = expand_selector_scopes(target)
        _require_requested(scopes, request.requested_mutation_scope)
        source_sheet, source_coord = parse_selector(source)
        target_sheet, target_range = parse_selector(target)
        if ":" in source_coord:
            raise SelectorError("formula source must be one cell")
        wb = load_workbook(path)
        formula = wb[source_sheet][source_coord].value
        if not isinstance(formula, str) or not formula.startswith("="):
            raise ValueError("source cell does not contain a formula")
        for cell in _cells(wb[target_sheet], target_range):
            cell.value = Translator(formula, origin=source_coord).translate_formula(cell.coordinate)
        wb.save(path); wb.close()
        return self._result(request, True, {"source": source, "filled": target}, scopes, (source, target))

    def _validate(self, request, path):
        try:
            wb = load_workbook(path, data_only=False, read_only=False)
            formula_errors = [{"selector": f"{ws.title}!{cell.coordinate}", "value": cell.value}
                              for ws in wb.worksheets for row in ws.iter_rows() for cell in row
                              if isinstance(cell.value, str) and cell.value.startswith("#")]
            output = {"valid": not formula_errors, "sheet_names": wb.sheetnames, "formula_errors": formula_errors}
            wb.close()
            return self._result(request, not formula_errors, output,
                                error="artifact_invalid:formula_error" if formula_errors else None)
        except Exception as error:
            return self._result(request, False, {"valid": False}, error=f"artifact_invalid:{type(error).__name__}:{error}")

    def _recalculate(self, request, path):
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        if not executable:
            return self._result(request, False, {"reason": "missing_soffice"}, error="capability_missing:soffice")
        # LibreOffice recalculates and rewrites the whole workbook.  Its process
        # success therefore does not establish that its effect stayed inside the
        # requested scope (for example, it may coerce unrelated date cells or
        # replace unsupported dynamic-array results with errors).  Preserve the
        # exact input bytes until the observed cell-level effect is authorized.
        original = path.read_bytes()
        before = _snapshot_cells(path)
        with tempfile.TemporaryDirectory() as directory:
            subprocess.run([executable, "--headless", "--convert-to", "xlsx", "--outdir", directory, str(path)],
                           check=True, capture_output=True, timeout=120)
            converted = Path(directory) / path.name
            if not converted.exists():
                raise OSError("recalculation produced no workbook")
            shutil.copy2(converted, path)
        after = _snapshot_cells(path)
        actual = tuple(Scope(key) for key in sorted(before.keys() | after.keys())
                       if before.get(key) != after.get(key))
        requested = request.requested_mutation_scope
        if not _scopes_authorized(actual, requested):
            path.write_bytes(original)
            outside = tuple(scope.resource for scope in actual
                            if not _scope_authorized(scope, requested))
            return self._result(
                request, False,
                {"recalculated": False, "rolled_back": True,
                 "outside_requested_scope": outside},
                actual,
                error="scope_violation:recalculation_changed_outside_requested_scope",
            )
        return self._result(request, True, {"recalculated": True}, actual)

    def _render(self, request, path):
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        if not executable:
            return self._result(request, False, {"reason": "missing_renderer"},
                                selectors=(request.inputs["selector"],), error="capability_missing:renderer")
        output_dir = Path(request.inputs["output_dir"]); output_dir.mkdir(parents=True, exist_ok=True)
        subprocess.run([executable, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
                       check=True, capture_output=True, timeout=120)
        rendered = output_dir / f"{path.stem}.pdf"
        if not rendered.exists():
            raise OSError("renderer produced no PDF")
        return self._result(request, True, {"rendered_path": str(rendered), "selector": request.inputs["selector"]},
                            selectors=(request.inputs["selector"],))


def _require_requested(actual: tuple[Scope, ...], requested: tuple[Scope, ...]) -> None:
    if not _scopes_authorized(actual, requested):
        raise ValueError("selector is outside requested mutation scope")


def _scope_authorized(actual: Scope, requested: tuple[Scope, ...]) -> bool:
    for allowed in requested:
        if allowed.resource == actual.resource:
            return True
        if allowed.resource.endswith("/*") and actual.resource.startswith(allowed.resource[:-1]):
            return True
    return False


def _scopes_authorized(actual: tuple[Scope, ...], requested: tuple[Scope, ...]) -> bool:
    return all(_scope_authorized(scope, requested) for scope in actual)


def manifests() -> tuple[CapabilityManifest, ...]:
    base_provenance = ("adapter_version", "artifact_sha256", "selectors")
    schemas = {
        "inspect_workbook": ({"selectors": {"type": "array"}}, ("selectors",), CapabilityEffect.OBSERVE),
        "write_cells": ({"writes": {"type": "array"}}, ("writes",), CapabilityEffect.MUTATE),
        "copy_or_fill_formula": ({"source": {"type": "string"}, "target": {"type": "string"}}, ("source", "target"), CapabilityEffect.MUTATE),
        "recalculate": ({}, (), CapabilityEffect.MUTATE),
        "validate_workbook": ({}, (), CapabilityEffect.VALIDATE),
        "render_range": ({"selector": {"type": "string"}, "output_dir": {"type": "string"}}, ("selector", "output_dir"), CapabilityEffect.RENDER),
    }
    return tuple(CapabilityManifest(name, VERSION, (KIND,),
        {"type": "object", "properties": props, "required": list(required), "additionalProperties": False},
        effect, (Scope("workbook/*"),) if effect is CapabilityEffect.MUTATE else (),
        provenance_requirements=base_provenance) for name, (props, required, effect) in schemas.items())


def register_spreadsheet_capabilities(broker) -> None:
    for manifest in manifests():
        broker.register(manifest, SpreadsheetCapability(manifest.name))
