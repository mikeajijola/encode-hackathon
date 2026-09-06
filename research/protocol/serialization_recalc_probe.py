"""Exercise actual recalculation and the broker in the canonical runtime."""
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook, load_workbook
from adapters.spreadsheet import KIND, VERSION, WorkbookSnapshots, register_spreadsheet_capabilities
from fulfilment import (Contract, DesiredAssertion, Scope, EvalSpec, CapabilityRequest,
                        Broker, EvidenceStore)

with TemporaryDirectory() as directory:
    artifact = Path(directory) / "probe.xlsx"
    wb = Workbook(); wb.active["A1"] = "=2+2"; wb.active["B1"] = "preserve"; wb.save(artifact); wb.close()
    scope = Scope("workbook/Sheet/A1")
    contract = Contract("probe", 1, "calculate A1", (DesiredAssertion("a", "A1 equals 4", (scope,)),),
                        ("preserve B1",), ("valid workbook",), (EvalSpec("e", "a", "independent"),), (scope,))
    broker = Broker(EvidenceStore(Path("/out/recalculation-probe.events.jsonl")), WorkbookSnapshots())
    register_spreadsheet_capabilities(broker)
    result = broker.invoke(contract, CapabilityRequest("probe", "recalculate", VERSION, str(artifact), KIND, {}, (scope,), ("d",)))
    wb = load_workbook(artifact, data_only=True)
    report = {"capability_succeeded": result.succeeded, "computed_value": wb.active["A1"].value,
              "preservation": wb.active["B1"].value == "preserve", "workbook_valid": True}
    wb.close()
    report["pass"] = result.succeeded and report["computed_value"] == 4 and report["preservation"]
    Path("/out/recalculation-probe.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report))
    if not report["pass"]:
        raise SystemExit(2)
