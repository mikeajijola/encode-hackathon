"""Offline retained-state strategy comparison; never imported by runtime."""
from __future__ import annotations

import argparse
from copy import copy
import difflib
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
from time import perf_counter
from zipfile import ZipFile

from openpyxl import load_workbook
from adapters.spreadsheet_diff import diff_workbooks
from adapters.spreadsheet_ooxml import patch_cells_from_candidate, patch_cell_values_raw
from fulfilment.models import Scope
from services.spreadsheet import _all_cells
from sb import load_dataset, load_answer_values, values_equal, recalculate
from evaluate import score_task


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def strict_changes(before, after):
    left, right = _all_cells(before), _all_cells(after)
    return [{"scope": key, "before": repr(left.get(key)), "after": repr(right.get(key))}
            for key in sorted(left.keys() | right.keys()) if left.get(key) != right.get(key)]


def run(dataset, retained, output, recalc=True):
    output.mkdir(parents=True, exist_ok=True)
    tasks = {task["id"]: task for task in load_dataset(dataset)}
    report = {"parent_commit": "c70e0b5", "kind": "offline_fixed_transition_replay",
              "model_calls": 0, "recalculation_enabled": recalc,
              "golden_access": "offline_scoring_only", "cases": []}
    for tid in ("58147", "61-4", "47766", "50971"):
        task = tasks[tid]
        original = Path(task["init_xlsx"])
        candidate = retained / "checkpoints" / tid / "mutation-01.xlsx"
        events = [json.loads(line) for line in (retained / "events" / f"{tid}.jsonl").read_text().splitlines()]
        contract = next(row["payload"]["contract"] for row in events
                        if row["event_type"] == "accepted_contract" and "contract" in row["payload"])
        scopes = tuple(Scope(s) for s in contract["authorized_mutation_scopes"])
        allowed = {s.resource for s in scopes}
        case_dir = output / tid; case_dir.mkdir(exist_ok=True)
        shutil.copy2(original, case_dir / "original.xlsx")
        shutil.copy2(candidate, case_dir / "retained-mutation-01.xlsx")
        case = {"task_id": tid, "original_sha256": digest(original),
                "retained_sha256": digest(candidate), "authorized_scope": sorted(allowed),
                "history": events, "strategies": []}
        for strategy in ("ordinary_openpyxl", "original_openpyxl_reconstruction", "ooxml_tree_transplant", "raw_xml_value_reconstruction"):
            target = case_dir / f"{strategy}.xlsx"
            start = perf_counter()
            if strategy == "ordinary_openpyxl":
                # The retained state is the output of the actual ordinary save.
                shutil.copy2(candidate, target)
                latency_kind = "retained_artifact_copy_only_not_original_save_latency"
            elif strategy == "original_openpyxl_reconstruction":
                wb, edited = load_workbook(original), load_workbook(candidate)
                for scope in scopes:
                    _, sheet, cell = scope.resource.split("/", 2)
                    wb[sheet][cell].value = copy(edited[sheet][cell].value)
                wb.save(target); wb.close(); edited.close()
                latency_kind = "reconstruction"
            else:
                function = patch_cells_from_candidate if strategy == "ooxml_tree_transplant" else patch_cell_values_raw
                function(original, candidate, target, scopes)
                latency_kind = "reconstruction"
            elapsed = (perf_counter() - start) * 1000
            try:
                changes = strict_changes(original, target)
                diff = diff_workbooks(original, target, scopes)
            except Exception as error:
                case["strategies"].append({"strategy": strategy, "artifact_valid": False,
                                          "error": f"{type(error).__name__}: {error}",
                                          "latency_ms": elapsed, "sha256": digest(target)})
                continue
            row = {"strategy": strategy, "latency_ms": elapsed, "latency_kind": latency_kind,
                   "tokens": 0, "additional_model_actions": 0, "replayed_mutations": 1,
                   "sha256": digest(target), "strict_changes": changes,
                   "outside_scope": [x for x in changes if x["scope"] not in allowed],
                   "diff": diff.evidence()}
            with ZipFile(original) as a, ZipFile(target) as b:
                xml_diffs = {}
                for name in diff.representation_parts:
                    if name.endswith((".xml", ".rels")):
                        left = a.read(name).decode() if name in a.namelist() else ""
                        right = b.read(name).decode() if name in b.namelist() else ""
                        xml_diffs[name] = list(difflib.unified_diff(left.replace("><", ">\n<").splitlines(),
                                                                right.replace("><", ">\n<").splitlines()))
                row["package_xml_diffs"] = xml_diffs
            with tempfile.TemporaryDirectory() as work:
                row["official"] = score_task(task, target, recalc, work)
            case["strategies"].append(row)
        # Compare candidate states through the exact official observable path,
        # after candidates have been generated. Never feeds runtime policy.
        with tempfile.TemporaryDirectory() as work:
            reference = recalculate(candidate, work) if recalc else candidate
            expected = load_answer_values(reference, task)
            for row in case["strategies"]:
                if row.get("artifact_valid") is False:
                    continue
                target = case_dir / f"{row['strategy']}.xlsx"
                with tempfile.TemporaryDirectory() as second:
                    visible = recalculate(target, second) if recalc else target
                    actual = load_answer_values(visible, task)
                row["official_path_difference_from_retained"] = [f"{k[0]}!{k[1]}" for k, value in expected.items()
                                                                 if not values_equal(value, actual.get(k))]
        report["cases"].append(case)
        (output / "strategy_comparison.json").write_text(json.dumps(report, indent=2, default=str) + "\n")
    (output / "strategy_comparison.json").write_text(json.dumps(report, indent=2, default=str) + "\n")
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", type=Path, required=True)
    parser.add_argument("--retained", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--no-recalc", action="store_true")
    args = parser.parse_args()
    result = run(args.dataset, args.retained, args.output, not args.no_recalc)
    print(json.dumps({"cases": len(result["cases"]), "model_calls": 0}))
